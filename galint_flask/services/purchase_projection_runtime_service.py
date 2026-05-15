from __future__ import annotations

from datetime import datetime, timedelta
from io import BytesIO
from math import isfinite
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import case, func, or_

from ..extensions import db
from ..models import (
    DocumentoEntradaEstoque,
    DocumentoEntradaEstoqueItem,
    FinanceLedgerEntry,
    FinanceSupplier,
    FinanceSupplierPreference,
    Item,
    StockBalance,
    StockMovement,
)
from .legacy_stock_normalizer import infer_packaging_measure, resolve_canonical_unit, resolve_packaging_factor
from .operation_visual_payload import OperationVisualPayloadService
from .potential_supplier_service import potential_supplier_service
from ..utils.report_branding import get_company_header_lines


class PurchaseProjectionService:
    VALID_BASE_UNITS = {"kg", "l", "m", "un"}
    REQUEST_UNIT_ALIASES = {"un", "unidade", "unidades", "peca", "pecas", "peça", "peças"}
    DEFAULT_WINDOW_DAYS = 30
    DEFAULT_COVERAGE_DAYS = 30
    MAX_WINDOW_DAYS = 365
    MAX_COVERAGE_DAYS = 365
    BALANCE_TOLERANCE = 1e-6

    STATUS_ALL = "all"
    STATUS_ACTIONABLE = "actionable"
    STATUS_CRITICAL = "critical"
    STATUS_ATTENTION = "attention"
    STATUS_NO_HISTORY = "no_history"
    STATUS_NO_CONSUMPTION = "no_consumption"
    STATUS_OK = "ok"

    FILTER_STATUS_OPTIONS = {
        STATUS_ALL,
        STATUS_ACTIONABLE,
        STATUS_CRITICAL,
        STATUS_ATTENTION,
        STATUS_NO_HISTORY,
        STATUS_NO_CONSUMPTION,
        STATUS_OK,
    }

    BLOCKING_VALIDATIONS = {"saldo_inconsistente", "unidade_invalida", "preco_incoerente"}

    ARCHITECTURE_COMPATIBILITY = {
        "is_fully_compatible": False,
        "implementation_mode": "safe_constraints",
        "answer": "Nao e totalmente compativel sem restricoes operacionais explicitas.",
        "safe_because": [
            "O GALINT ja possui ledger por StockMovement, cache por StockBalance, conversao por unidade base e trilha documental/financeira para preco e fornecedor.",
            "A pagina pode operar com seguranca quando usa saldo e consumo apenas do ledger, mantem a unidade base como contrato e explicita lacunas de fornecedor, preco e historico.",
        ],
        "conflicts": [
            "Nem todos os itens estao no mesmo estagio de reconciliacao entre legado e ledger, entao nao e seguro chamar a funcionalidade de totalmente irrestrita.",
            "Fornecedor e preco nao estao completos de forma uniforme para todos os itens; misturar origens incoerentes repetiria inconsistencias ja vistas na base.",
        ],
        "required_adaptations": [
            "Usar StockMovement como fonte autoritativa de saldo e consumo, deixando StockBalance apenas como detector de drift.",
            "Calcular tudo somente em unidade base valida do GALINT e nao usar embalagem como saldo autoritativo.",
            "Anexar fornecedor e preco apenas quando houver par coerente ou declarar explicitamente a ausencia no payload e na tela.",
        ],
        "risk_of_repeating_existing_inconsistencies": [
            "Usar cache ou embalagem como verdade de saldo reabre divergencia com o ledger.",
            "Cruzar fornecedor de uma origem com preco de outra origem sem compatibilidade explicita pode montar pedido errado.",
        ],
    }

    VALIDATION_MESSAGES = {
        "saldo_inconsistente": "Saldo inconsistente entre ledger e cache StockBalance.",
        "unidade_invalida": "Item fora das unidades base validas do GALINT (kg, l, m, un).",
        "preco_incoerente": "Preco base ausente, zero ou invalido para compra.",
        "item_sem_consumo": "Item sem consumo no periodo analisado.",
        "item_sem_historico": "Item sem historico de consumo no ledger.",
        "item_sem_fornecedor": "Item sem fornecedor coerente para consolidacao do pedido.",
        "quantidade_manual_invalida": "Quantidade manual invalida; a sugestao calculada foi mantida.",
        "quantidade_manual_zerada": "Quantidade manual zerada ou negativa; o item nao entrou no carrinho.",
    }

    STATUS_LABELS = {
        STATUS_CRITICAL: "Critico",
        STATUS_ATTENTION: "Atencao",
        STATUS_NO_HISTORY: "Sem historico",
        STATUS_NO_CONSUMPTION: "Sem consumo",
        STATUS_OK: "Ok",
    }

    STATUS_BADGE_CLASSES = {
        STATUS_CRITICAL: "danger",
        STATUS_ATTENTION: "warning",
        STATUS_NO_HISTORY: "secondary",
        STATUS_NO_CONSUMPTION: "secondary",
        STATUS_OK: "success",
    }

    PRICE_SOURCE_LABELS = {
        "document_pair": "Documento fiscal coerente",
        "finance_pair": "Financeiro coerente",
        "preferred_item_master": "Fornecedor preferencial + preco mestre",
        "document_generic": "Documento sem fornecedor coerente",
        "item_master": "Preco mestre do item",
        "none": "Sem preco coerente",
    }

    @classmethod
    def get_architecture_compatibility(cls) -> dict[str, Any]:
        return {
            **cls.ARCHITECTURE_COMPATIBILITY,
            "checked_at": datetime.utcnow().isoformat(),
        }

    @classmethod
    def normalize_filters(
        cls,
        *,
        window_days: object = None,
        coverage_days: object = None,
        search: object = None,
        category: object = None,
        brand: object = None,
        status: object = None,
        include_inactive: object = None,
    ) -> dict[str, Any]:
        return {
            "window_days": cls._normalize_int(window_days, default=cls.DEFAULT_WINDOW_DAYS, minimum=1, maximum=cls.MAX_WINDOW_DAYS),
            "coverage_days": cls._normalize_int(coverage_days, default=cls.DEFAULT_COVERAGE_DAYS, minimum=1, maximum=cls.MAX_COVERAGE_DAYS),
            "search": str(search or "").strip(),
            "category": str(category or "").strip(),
            "brand": str(brand or "").strip(),
            "status": cls._normalize_status(status),
            "include_inactive": cls._normalize_bool(include_inactive),
        }

    @classmethod
    def build_filter_options(cls, *, include_inactive: object = None) -> dict[str, list[str]]:
        base_query = Item.query
        if not cls._normalize_bool(include_inactive) and hasattr(Item, "ativo"):
            base_query = base_query.filter(Item.ativo.is_(True))

        categories = sorted(
            {
                str(value or "").strip()
                for (value,) in base_query.with_entities(Item.categoria).distinct().all()
                if str(value or "").strip()
            },
            key=str.casefold,
        )
        brands = sorted(
            {
                str(value or "").strip()
                for (value,) in base_query.with_entities(Item.marca).distinct().all()
                if str(value or "").strip()
            },
            key=str.casefold,
        )

        return {
            "categories": categories,
            "brands": brands,
        }

    @classmethod
    def build_projection_report(
        cls,
        *,
        window_days: object = None,
        coverage_days: object = None,
        search: object = None,
        category: object = None,
        brand: object = None,
        status: object = None,
        include_inactive: object = None,
        selected_codes: list[str] | tuple[str, ...] | set[str] | None = None,
        manual_quantities: dict[str, object] | None = None,
        potential_quote_ids: dict[str, object] | None = None,
    ) -> dict[str, Any]:
        filters = cls.normalize_filters(
            window_days=window_days,
            coverage_days=coverage_days,
            search=search,
            category=category,
            brand=brand,
            status=status,
            include_inactive=include_inactive,
        )
        filter_options = cls.build_filter_options(include_inactive=filters.get("include_inactive"))
        selected_set = {str(code or "").strip() for code in (selected_codes or []) if str(code or "").strip()}
        items = cls._load_items(filters)
        filtered_codes = {str(item.codigo_item or "").strip() for item in items if str(item.codigo_item or "").strip()}
        supplemental_items = cls._load_items_by_codes(selected_set - filtered_codes)
        item_map: dict[str, Item] = {}
        for item in [*items, *supplemental_items]:
            codigo_item = str(item.codigo_item or "").strip()
            if not codigo_item or codigo_item in item_map:
                continue
            item_map[codigo_item] = item

        product_ids = list(item_map.keys())
        ledger_balances = cls._load_ledger_balances(product_ids)
        cached_balances = cls._load_cached_balances(product_ids)
        window_consumption = cls._load_window_consumption(product_ids, window_days=int(filters["window_days"]))
        consumption_history = cls._load_consumption_history(product_ids)
        preferences = cls._load_preferences(product_ids)
        document_candidates = cls._load_latest_document_candidates(product_ids)
        finance_candidates = cls._load_latest_finance_candidates(product_ids)
        supplier_lookup = cls._load_supplier_lookup(
            {
                *[pref.fornecedor_id for pref in preferences.values() if getattr(pref, "fornecedor_id", None)],
                *[
                    (candidate.get("supplier") or {}).get("id")
                    for candidate in document_candidates.values()
                    if (candidate.get("supplier") or {}).get("id")
                ],
                *[
                    (candidate.get("supplier") or {}).get("id")
                    for candidate in finance_candidates.values()
                    if (candidate.get("supplier") or {}).get("id")
                ],
            }
        )

        all_rows = [
            cls._build_row(
                item,
                filters=filters,
                ledger_balances=ledger_balances,
                cached_balances=cached_balances,
                window_consumption=window_consumption,
                consumption_history=consumption_history,
                preferences=preferences,
                document_candidates=document_candidates,
                finance_candidates=finance_candidates,
                supplier_lookup=supplier_lookup,
            )
            for item in item_map.values()
        ]
        reference_prices = {
            str(row.get("codigo_item") or "").strip(): row.get("price_unit_base")
            for row in all_rows
            if str(row.get("codigo_item") or "").strip()
        }
        potential_quotes = potential_supplier_service.list_best_quotes_by_item(
            product_ids,
            limit_per_item=5,
            reference_prices=reference_prices,
        )
        cls._attach_potential_quotes(all_rows, potential_quotes, potential_quote_ids=potential_quote_ids)
        row_lookup = {str(row.get("codigo_item") or "").strip(): row for row in all_rows if str(row.get("codigo_item") or "").strip()}
        rows = [row_lookup[codigo_item] for codigo_item in filtered_codes if codigo_item in row_lookup]
        visible_rows = [row for row in rows if cls._status_matches_filter(str(filters["status"]), row)]
        visible_rows.sort(key=cls._row_sort_key)
        cart = cls._apply_selection(
            all_rows,
            selected_codes=selected_codes,
            manual_quantities=manual_quantities,
            potential_quote_ids=potential_quote_ids,
        )
        return {
            "compatibility": cls.get_architecture_compatibility(),
            "generated_at": datetime.utcnow().isoformat(),
            "filters": filters,
            "filter_options": filter_options,
            "summary": cls._build_summary(rows, visible_rows, cart),
            "rows": visible_rows,
            "cart": cart,
            "visible_categories": cls._build_visible_category_groups(visible_rows),
        }

    @classmethod
    def build_workbook(cls, report: dict[str, Any], *, requested_by: dict[str, Any] | None = None) -> BytesIO:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Pedido"
        sheet.sheet_view.showGridLines = False
        filters = dict(report.get("filters") or {})
        cart = dict(report.get("cart") or {})
        requested_by = dict(requested_by or {})

        total_columns = 19
        end_column = get_column_letter(total_columns)
        border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )
        unlocked_protection = Protection(locked=False)
        locked_protection = Protection(locked=True)
        title_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        header_fill = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")
        total_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
        footer_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
        value_fill = PatternFill(start_color="FDE9D9", end_color="FDE9D9", fill_type="solid")

        current_row = 1
        company_lines = get_company_header_lines()
        for index, line in enumerate(company_lines):
            sheet.merge_cells(f"A{current_row}:{end_column}{current_row}")
            cell = sheet.cell(row=current_row, column=1)
            cell.value = line
            cell.font = Font(bold=(index == 0), size=18 if index == 0 else 11, color="0F172A")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.protection = unlocked_protection
            sheet.row_dimensions[current_row].height = 28 if index == 0 else 18
            current_row += 1

        sheet.merge_cells(f"A{current_row}:{end_column}{current_row}")
        title_cell = sheet.cell(row=current_row, column=1)
        title_cell.value = "PROJECAO DE COMPRAS"
        title_cell.fill = title_fill
        title_cell.font = Font(bold=True, size=16, color="FFFFFF")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        title_cell.border = border
        title_cell.protection = unlocked_protection
        sheet.row_dimensions[current_row].height = 26
        current_row += 2

        requester_label = str(requested_by.get("nome") or "Solicitante nao informado").strip() or "Solicitante nao informado"
        requested_at_label = str(requested_by.get("requested_at") or report.get("generated_at") or "").strip()

        headers = [
            "Fornecedor Fiscal",
            "CNPJ Fiscal",
            "Descricao",
            "Marca",
            "Categoria",
            "Pedido",
            "Preco Atual",
            "Total Atual",
            "Potencial Fornecedor",
            "Link Fornecedor",
            "CNPJ Potencial",
            "Link Produto",
            "Preco Potencial",
            "Total Potencial",
            "Variacao",
            "Economia Potencial",
            "Fonte Cotacao",
            "Contato",
            "Endereco",
        ]
        header_row = current_row
        for column_index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=header_row, column=column_index, value=header)
            cell.fill = header_fill
            cell.font = Font(bold=True, size=12, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
            cell.protection = locked_protection if column_index in {6, 7, 8, 13, 14, 16} else unlocked_protection
        sheet.row_dimensions[header_row].height = 24
        sheet.freeze_panes = f"A{header_row + 1}"

        data_start_row = header_row + 1
        category_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
        category_total_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        export_categories = list(cart.get("category_groups") or [])

        if not export_categories:
            sheet.merge_cells(f"A{data_start_row}:{end_column}{data_start_row}")
            cell = sheet.cell(row=data_start_row, column=1)
            cell.value = "Nenhum item selecionado para exportacao."
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = Font(size=12, italic=True, color="64748B")
            cell.protection = unlocked_protection
            data_end_row = data_start_row
        else:
            current_data_row = data_start_row
            for category_index, category_group in enumerate(export_categories, start=1):
                category_name = str(category_group.get("category_name") or "Sem categoria").strip() or "Sem categoria"
                items = list(category_group.get("items") or [])
                if not items:
                    continue

                sheet.merge_cells(start_row=current_data_row, start_column=1, end_row=current_data_row, end_column=total_columns)
                category_cell = sheet.cell(row=current_data_row, column=1)
                category_cell.value = f"CATEGORIA: {category_name}"
                category_cell.fill = category_fill
                category_cell.font = Font(bold=True, size=12, color="1E3A8A")
                category_cell.alignment = Alignment(horizontal="left", vertical="center")
                category_cell.border = border
                category_cell.protection = unlocked_protection
                sheet.row_dimensions[current_data_row].height = 22
                current_data_row += 1

                for item in items:
                    supplier = dict(item.get("supplier") or {})
                    potential_quote = dict(item.get("selected_potential_quote") or {})
                    potential_supplier = dict(potential_quote.get("supplier") or {})
                    supplier_link = str(
                        potential_supplier.get("contact_url")
                        or potential_supplier.get("source_url")
                        or potential_supplier.get("website")
                        or potential_quote.get("product_url")
                        or ""
                    ).strip()
                    row_index = current_data_row
                    values = [
                        str(supplier.get("name") or "Sem fornecedor identificado").strip() or "Sem fornecedor identificado",
                        str(supplier.get("cnpj") or "").strip(),
                        str(item.get("descricao") or "").strip(),
                        str(item.get("marca") or "Sem marca").strip() or "Sem marca",
                        category_name,
                        cls._format_number_input_value(item.get("requested_quantity_input") or 0.0),
                        item.get("price_unit_request") if cls._is_positive_number(item.get("price_unit_request")) else item.get("price_unit_base"),
                        item.get("requested_total_value"),
                        str(potential_supplier.get("display_name") or "").strip(),
                        supplier_link,
                        str(potential_supplier.get("cnpj") or "").strip(),
                        str(potential_quote.get("product_url") or "").strip(),
                        potential_quote.get("unit_price_base"),
                        potential_quote.get("requested_total_value"),
                        str(potential_quote.get("variation_display") or "").strip(),
                        potential_quote.get("estimated_savings_total"),
                        str(potential_quote.get("source_name") or "").strip(),
                        str(potential_supplier.get("contact_url") or potential_quote.get("product_url") or "").strip(),
                        str(potential_supplier.get("address_display") or "").strip(),
                    ]
                    for column_index, value in enumerate(values, start=1):
                        cell = sheet.cell(row=row_index, column=column_index, value=value)
                        cell.border = border
                        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                        cell.font = Font(size=12, bold=(column_index == 3), color="0F172A")
                        cell.protection = unlocked_protection
                        if column_index == 6:
                            cell.alignment = Alignment(horizontal="center", vertical="center")
                            cell.protection = locked_protection
                            cell.font = Font(size=12, bold=True, color="0F172A")
                        if column_index in {7, 8, 13, 14, 16}:
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                            cell.protection = locked_protection
                        if column_index in {7, 8, 13, 14, 16} and isinstance(value, (int, float)):
                            cell.number_format = 'R$ #,##0.00'
                            cell.fill = value_fill
                            cell.font = Font(size=12, italic=True, color="0F172A")
                        if column_index in {10, 12, 18} and value:
                            cell.hyperlink = str(value)
                            cell.style = "Hyperlink"
                    if row_index % 2 == 0:
                        for column_index in range(1, total_columns + 1):
                            if column_index in {7, 8, 13, 14, 16}:
                                continue
                            sheet.cell(row=row_index, column=column_index).fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
                    sheet.row_dimensions[row_index].height = 24
                    current_data_row += 1

                subtotal_row = current_data_row
                sheet.merge_cells(start_row=subtotal_row, start_column=1, end_row=subtotal_row, end_column=5)
                subtotal_label_cell = sheet.cell(row=subtotal_row, column=1)
                subtotal_label_cell.value = f"SUBTOTAL DA CATEGORIA: {category_name}"
                subtotal_label_cell.fill = category_total_fill
                subtotal_label_cell.font = Font(bold=True, size=12, color="166534")
                subtotal_label_cell.alignment = Alignment(horizontal="right", vertical="center")
                subtotal_label_cell.border = border
                subtotal_label_cell.protection = unlocked_protection

                subtotal_quantity_cell = sheet.cell(
                    row=subtotal_row,
                    column=6,
                    value=cls._format_number_input_value(
                        sum(float(item.get("requested_quantity_input") or 0.0) for item in items)
                    ),
                )
                subtotal_quantity_cell.fill = category_total_fill
                subtotal_quantity_cell.font = Font(bold=True, size=12, color="166534")
                subtotal_quantity_cell.alignment = Alignment(horizontal="center", vertical="center")
                subtotal_quantity_cell.border = border
                subtotal_quantity_cell.protection = locked_protection

                subtotal_price_cell = sheet.cell(row=subtotal_row, column=7, value="")
                subtotal_price_cell.fill = category_total_fill
                subtotal_price_cell.border = border
                subtotal_price_cell.protection = locked_protection

                subtotal_value_cell = sheet.cell(
                    row=subtotal_row,
                    column=8,
                    value=float(category_group.get("requested_value_total") or 0.0),
                )
                subtotal_value_cell.fill = category_total_fill
                subtotal_value_cell.font = Font(bold=True, size=12, color="166534")
                subtotal_value_cell.alignment = Alignment(horizontal="right", vertical="center")
                subtotal_value_cell.border = border
                subtotal_value_cell.protection = locked_protection
                subtotal_value_cell.number_format = 'R$ #,##0.00'

                subtotal_potential_value = sum(
                    float(((item.get("selected_potential_quote") or {}).get("requested_total_value")) or 0.0)
                    for item in items
                )
                subtotal_savings_value = sum(
                    float(((item.get("selected_potential_quote") or {}).get("estimated_savings_total")) or 0.0)
                    for item in items
                )
                for column_index in range(9, total_columns + 1):
                    value = ""
                    if column_index == 14:
                        value = subtotal_potential_value
                    elif column_index == 16:
                        value = subtotal_savings_value
                    cell = sheet.cell(row=subtotal_row, column=column_index, value=value)
                    cell.fill = category_total_fill
                    cell.border = border
                    cell.protection = locked_protection
                    if column_index in {14, 16}:
                        cell.font = Font(bold=True, size=12, color="166534")
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        cell.number_format = 'R$ #,##0.00'

                current_data_row += 1
                if category_index < len(export_categories):
                    current_data_row += 1

            data_end_row = current_data_row - 1

        total_row = data_end_row + 1
        sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=5)
        total_label_cell = sheet.cell(row=total_row, column=1)
        total_label_cell.value = "SUBTOTAL GERAL"
        total_label_cell.fill = total_fill
        total_label_cell.font = Font(bold=True, size=12, color="7C2D12")
        total_label_cell.alignment = Alignment(horizontal="right", vertical="center")
        total_label_cell.border = border
        total_label_cell.protection = unlocked_protection

        total_formulas = {
            6: cls._format_number_input_value(
                sum(float(item.get("requested_quantity_input") or 0.0) for group in export_categories for item in (group.get("items") or []))
            ) if export_categories else "0",
            7: "",
            8: float(cart.get("requested_value_total") or 0.0) if export_categories else 0,
            9: "",
            10: "",
            11: "",
            12: "",
            13: "",
            14: sum(
                float(((item.get("selected_potential_quote") or {}).get("requested_total_value")) or 0.0)
                for group in export_categories
                for item in (group.get("items") or [])
            ) if export_categories else 0,
            15: "",
            16: sum(
                float(((item.get("selected_potential_quote") or {}).get("estimated_savings_total")) or 0.0)
                for group in export_categories
                for item in (group.get("items") or [])
            ) if export_categories else 0,
            17: "",
            18: "",
            19: "",
        }
        for column_index in range(6, total_columns + 1):
            cell = sheet.cell(row=total_row, column=column_index, value=total_formulas[column_index])
            cell.fill = total_fill if column_index == 6 else value_fill
            cell.font = Font(bold=True, size=16, color="0F172A") if column_index == 6 else Font(size=12, italic=True, color="0F172A")
            cell.alignment = Alignment(horizontal="center", vertical="center") if column_index == 6 else Alignment(horizontal="right", vertical="center")
            cell.border = border
            cell.protection = locked_protection
            if column_index in {8, 14, 16}:
                cell.number_format = 'R$ #,##0.00'

        footer_row = total_row + 3
        sheet.merge_cells(f"A{footer_row}:{end_column}{footer_row}")
        footer_cell = sheet.cell(row=footer_row, column=1)
        footer_cell.value = f"Emitido por {requester_label} em {requested_at_label}"
        footer_cell.fill = footer_fill
        footer_cell.font = Font(bold=True, size=12, color="0F172A")
        footer_cell.alignment = Alignment(horizontal="center", vertical="center")
        footer_cell.border = border
        footer_cell.protection = unlocked_protection
        sheet.row_dimensions[footer_row].height = 24

        sheet.column_dimensions["A"].width = 28
        sheet.column_dimensions["B"].width = 22
        sheet.column_dimensions["C"].width = 44
        sheet.column_dimensions["D"].width = 20
        sheet.column_dimensions["E"].width = 24
        sheet.column_dimensions["F"].width = 16
        sheet.column_dimensions["G"].width = 20
        sheet.column_dimensions["H"].width = 18
        sheet.column_dimensions["I"].width = 30
        sheet.column_dimensions["J"].width = 42
        sheet.column_dimensions["K"].width = 20
        sheet.column_dimensions["L"].width = 42
        sheet.column_dimensions["M"].width = 18
        sheet.column_dimensions["N"].width = 18
        sheet.column_dimensions["O"].width = 14
        sheet.column_dimensions["P"].width = 20
        sheet.column_dimensions["Q"].width = 20
        sheet.column_dimensions["R"].width = 42
        sheet.column_dimensions["S"].width = 34
        sheet.page_setup.orientation = "landscape"
        sheet.protection.sheet = True
        sheet.protection.enable()
        sheet.protection.formatCells = False
        sheet.protection.formatColumns = False
        sheet.protection.formatRows = False
        sheet.protection.insertColumns = False
        sheet.protection.insertRows = False
        sheet.protection.deleteColumns = False
        sheet.protection.deleteRows = False

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer

    @classmethod
    def _load_items(cls, filters: dict[str, Any]) -> list[Item]:
        query = Item.query
        if not filters.get("include_inactive") and hasattr(Item, "ativo"):
            query = query.filter(Item.ativo.is_(True))
        category = str(filters.get("category") or "").strip()
        if category:
            query = query.filter(func.lower(func.coalesce(Item.categoria, "")) == category.casefold())
        brand = str(filters.get("brand") or "").strip()
        if brand:
            query = query.filter(func.lower(func.coalesce(Item.marca, "")) == brand.casefold())
        search = str(filters.get("search") or "").strip()
        if search:
            like = f"%{search}%"
            query = query.filter(
                or_(
                    Item.codigo_item.ilike(like),
                    Item.descricao.ilike(like),
                    Item.categoria.ilike(like),
                    Item.marca.ilike(like),
                )
            )
        return query.order_by(Item.descricao.asc(), Item.codigo_item.asc()).all()

    @classmethod
    def _load_items_by_codes(cls, product_ids: set[str] | list[str] | tuple[str, ...]) -> list[Item]:
        normalized_ids = [str(product_id or "").strip() for product_id in product_ids if str(product_id or "").strip()]
        if not normalized_ids:
            return []
        rows = Item.query.filter(Item.codigo_item.in_(normalized_ids)).all()
        rows.sort(key=lambda item: normalized_ids.index(str(item.codigo_item or "").strip()))
        return rows

    @classmethod
    def _load_ledger_balances(cls, product_ids: list[str]) -> dict[str, float]:
        if not product_ids:
            return {}
        rows = (
            db.session.query(
                StockMovement.product_id,
                func.coalesce(func.sum(StockMovement.quantity_base), 0.0),
            )
            .filter(StockMovement.product_id.in_(product_ids))
            .group_by(StockMovement.product_id)
            .all()
        )
        return {str(product_id): float(total or 0.0) for product_id, total in rows if product_id}

    @classmethod
    def _load_cached_balances(cls, product_ids: list[str]) -> dict[str, float]:
        if not product_ids:
            return {}
        rows = StockBalance.query.filter(StockBalance.product_id.in_(product_ids)).all()
        return {str(row.product_id): float(row.quantity_base or 0.0) for row in rows if row.product_id}

    @classmethod
    def _load_window_consumption(cls, product_ids: list[str], *, window_days: int) -> dict[str, dict[str, Any]]:
        if not product_ids:
            return {}
        cutoff = datetime.utcnow() - timedelta(days=window_days)
        rows = (
            db.session.query(
                StockMovement.product_id,
                func.coalesce(
                    func.sum(
                        case(
                            (StockMovement.movement_type == "saida", -StockMovement.quantity_base),
                            else_=0.0,
                        )
                    ),
                    0.0,
                ).label("saida_total"),
                func.coalesce(
                    func.sum(
                        case(
                            (StockMovement.movement_type == "devolucao", StockMovement.quantity_base),
                            else_=0.0,
                        )
                    ),
                    0.0,
                ).label("devolucao_total"),
                func.max(
                    case(
                        (StockMovement.movement_type == "saida", StockMovement.created_at),
                        else_=None,
                    )
                ).label("last_saida_at"),
            )
            .filter(StockMovement.product_id.in_(product_ids))
            .filter(StockMovement.created_at >= cutoff)
            .filter(StockMovement.movement_type.in_(["saida", "devolucao"]))
            .group_by(StockMovement.product_id)
            .all()
        )
        payload: dict[str, dict[str, Any]] = {}
        for product_id, saida_total, devolucao_total, last_saida_at in rows:
            if not product_id:
                continue
            payload[str(product_id)] = {
                "consumption_window_base": max(float(saida_total or 0.0) - float(devolucao_total or 0.0), 0.0),
                "last_saida_at": last_saida_at.isoformat() if last_saida_at else None,
            }
        return payload

    @classmethod
    def _load_consumption_history(cls, product_ids: list[str]) -> dict[str, dict[str, Any]]:
        if not product_ids:
            return {}
        rows = (
            db.session.query(
                StockMovement.product_id,
                func.sum(case((StockMovement.movement_type == "saida", 1), else_=0)).label("saida_count"),
                func.max(case((StockMovement.movement_type == "saida", StockMovement.created_at), else_=None)).label("last_saida_at"),
            )
            .filter(StockMovement.product_id.in_(product_ids))
            .filter(StockMovement.movement_type.in_(["saida", "devolucao"]))
            .group_by(StockMovement.product_id)
            .all()
        )
        return {
            str(product_id): {
                "saida_count": int(saida_count or 0),
                "last_saida_at": last_saida_at.isoformat() if last_saida_at else None,
            }
            for product_id, saida_count, last_saida_at in rows
            if product_id
        }

    @classmethod
    def _load_preferences(cls, product_ids: list[str]) -> dict[str, FinanceSupplierPreference]:
        if not product_ids:
            return {}
        rows = FinanceSupplierPreference.query.filter(FinanceSupplierPreference.codigo_item.in_(product_ids)).all()
        return {str(row.codigo_item): row for row in rows if row.codigo_item}

    @classmethod
    def _load_latest_document_candidates(cls, product_ids: list[str]) -> dict[str, dict[str, Any]]:
        if not product_ids:
            return {}
        rows = (
            db.session.query(
                DocumentoEntradaEstoqueItem.codigo_item,
                DocumentoEntradaEstoqueItem.valor_unitario_base,
                DocumentoEntradaEstoqueItem.id_documento_item,
                DocumentoEntradaEstoque.numero_documento,
                DocumentoEntradaEstoque.tipo_documento,
                DocumentoEntradaEstoque.data_recebimento,
                DocumentoEntradaEstoque.data_emissao,
                DocumentoEntradaEstoque.fornecedor_id,
                DocumentoEntradaEstoque.fornecedor_nome,
                DocumentoEntradaEstoque.cnpj_emitente,
            )
            .join(DocumentoEntradaEstoque, DocumentoEntradaEstoque.id_documento == DocumentoEntradaEstoqueItem.documento_id)
            .filter(DocumentoEntradaEstoqueItem.codigo_item.in_(product_ids))
            .filter(DocumentoEntradaEstoqueItem.valor_unitario_base.isnot(None))
            .filter(DocumentoEntradaEstoqueItem.valor_unitario_base > 0)
            .order_by(
                DocumentoEntradaEstoque.data_recebimento.desc(),
                DocumentoEntradaEstoque.data_emissao.desc(),
                DocumentoEntradaEstoqueItem.id_documento_item.desc(),
            )
            .all()
        )
        result: dict[str, dict[str, Any]] = {}
        for row in rows:
            codigo = str(row.codigo_item or "").strip()
            if not codigo or codigo in result:
                continue
            result[codigo] = {
                "unit_price_base": float(row.valor_unitario_base or 0.0),
                "document_number": row.numero_documento,
                "document_type": row.tipo_documento,
                "document_date": row.data_recebimento.isoformat() if row.data_recebimento else (row.data_emissao.isoformat() if row.data_emissao else None),
                "supplier": {
                    "id": int(row.fornecedor_id) if row.fornecedor_id else None,
                    "name": (row.fornecedor_nome or "").strip() or None,
                    "cnpj": (row.cnpj_emitente or "").strip() or None,
                },
            }
        return result

    @classmethod
    def _load_latest_finance_candidates(cls, product_ids: list[str]) -> dict[str, dict[str, Any]]:
        if not product_ids:
            return {}
        rows = (
            db.session.query(
                FinanceLedgerEntry.codigo_item,
                FinanceLedgerEntry.valor_unitario_base,
                FinanceLedgerEntry.numero_documento,
                FinanceLedgerEntry.tipo_documento,
                FinanceLedgerEntry.data_recebimento_documento,
                FinanceLedgerEntry.data_emissao_documento,
                FinanceLedgerEntry.data_lancamento,
                FinanceLedgerEntry.fornecedor_id,
                FinanceSupplier.nome_fantasia,
                FinanceSupplier.razao_social,
                FinanceSupplier.cnpj,
            )
            .outerjoin(FinanceSupplier, FinanceSupplier.id == FinanceLedgerEntry.fornecedor_id)
            .filter(FinanceLedgerEntry.codigo_item.in_(product_ids))
            .filter(FinanceLedgerEntry.valor_unitario_base.isnot(None))
            .filter(FinanceLedgerEntry.valor_unitario_base > 0)
            .order_by(FinanceLedgerEntry.data_lancamento.desc(), FinanceLedgerEntry.id.desc())
            .all()
        )
        result: dict[str, dict[str, Any]] = {}
        for row in rows:
            codigo = str(row.codigo_item or "").strip()
            if not codigo or codigo in result:
                continue
            result[codigo] = {
                "unit_price_base": float(row.valor_unitario_base or 0.0),
                "document_number": row.numero_documento,
                "document_type": row.tipo_documento,
                "document_date": row.data_recebimento_documento.isoformat() if row.data_recebimento_documento else (row.data_emissao_documento.isoformat() if row.data_emissao_documento else (row.data_lancamento.isoformat() if row.data_lancamento else None)),
                "supplier": {
                    "id": int(row.fornecedor_id) if row.fornecedor_id else None,
                    "name": (row.nome_fantasia or row.razao_social or "").strip() or None,
                    "cnpj": (row.cnpj or "").strip() or None,
                },
            }
        return result

    @classmethod
    def _load_supplier_lookup(cls, supplier_ids: set[object]) -> dict[int, FinanceSupplier]:
        normalized_ids = {int(supplier_id) for supplier_id in supplier_ids if supplier_id}
        if not normalized_ids:
            return {}
        rows = FinanceSupplier.query.filter(FinanceSupplier.id.in_(sorted(normalized_ids))).all()
        return {int(row.id): row for row in rows if row.id is not None}

    @classmethod
    def _build_row(
        cls,
        item: Item,
        *,
        filters: dict[str, Any],
        ledger_balances: dict[str, float],
        cached_balances: dict[str, float],
        window_consumption: dict[str, dict[str, Any]],
        consumption_history: dict[str, dict[str, Any]],
        preferences: dict[str, FinanceSupplierPreference],
        document_candidates: dict[str, dict[str, Any]],
        finance_candidates: dict[str, dict[str, Any]],
        supplier_lookup: dict[int, FinanceSupplier],
    ) -> dict[str, Any]:
        codigo = str(item.codigo_item or "").strip()
        unit_base = cls._normalize_unit(resolve_canonical_unit(item))
        balance_base = float(ledger_balances.get(codigo, 0.0) or 0.0)
        cached_balance = cached_balances.get(codigo)
        window_data = window_consumption.get(codigo) or {}
        history_data = consumption_history.get(codigo) or {}
        procurement = cls._resolve_procurement(
            item,
            preference=preferences.get(codigo),
            document_candidate=document_candidates.get(codigo),
            finance_candidate=finance_candidates.get(codigo),
            supplier_lookup=supplier_lookup,
        )

        validation_codes: list[str] = []
        if unit_base not in cls.VALID_BASE_UNITS:
            validation_codes.append("unidade_invalida")
        if cached_balance is not None and abs(float(cached_balance or 0.0) - balance_base) > cls.BALANCE_TOLERANCE:
            validation_codes.append("saldo_inconsistente")
        if int(history_data.get("saida_count") or 0) <= 0:
            validation_codes.append("item_sem_historico")
        consumption_window_base = max(float(window_data.get("consumption_window_base") or 0.0), 0.0)
        if int(history_data.get("saida_count") or 0) > 0 and consumption_window_base <= cls.BALANCE_TOLERANCE:
            validation_codes.append("item_sem_consumo")
        if not cls._is_positive_number((procurement.get("price") or {}).get("unit_price_base")):
            validation_codes.append("preco_incoerente")
        if not str((procurement.get("supplier") or {}).get("name") or "").strip():
            validation_codes.append("item_sem_fornecedor")

        window_days = int(filters.get("window_days") or cls.DEFAULT_WINDOW_DAYS)
        coverage_days = int(filters.get("coverage_days") or cls.DEFAULT_COVERAGE_DAYS)
        consumption_average_base = consumption_window_base / float(window_days) if consumption_window_base > cls.BALANCE_TOLERANCE else 0.0
        days_remaining = (balance_base / consumption_average_base) if consumption_average_base > cls.BALANCE_TOLERANCE else None
        suggested_quantity_base = max((consumption_average_base * float(coverage_days)) - balance_base, 0.0) if consumption_average_base > cls.BALANCE_TOLERANCE else 0.0
        price_unit_base = (procurement.get("price") or {}).get("unit_price_base")
        estimated_suggestion_value = round(float(price_unit_base) * suggested_quantity_base, 2) if cls._is_positive_number(price_unit_base) and suggested_quantity_base > cls.BALANCE_TOLERANCE else None
        request_quantity_contract = cls._build_request_quantity_contract(item, unit_base)
        request_factor_base = float(request_quantity_contract.get("factor_base") or 1.0)
        suggested_quantity_input = suggested_quantity_base / request_factor_base if request_factor_base > cls.BALANCE_TOLERANCE else suggested_quantity_base
        request_unit_price = round(float(price_unit_base) * request_factor_base, 4) if cls._is_positive_number(price_unit_base) else None
        status = cls._classify_row_status(
            balance_base=balance_base,
            consumption_average_base=consumption_average_base,
            suggested_quantity_base=suggested_quantity_base,
            validation_codes=validation_codes,
        )
        validation_notes = [
            {
                "code": code,
                "message": cls.VALIDATION_MESSAGES.get(code, code),
                "blocking": code in cls.BLOCKING_VALIDATIONS,
            }
            for code in validation_codes
        ]

        return {
            "codigo_item": codigo,
            "descricao": item.descricao,
            "categoria": item.categoria,
            "marca": item.marca,
            "ativo": bool(getattr(item, "ativo", True)),
            "unit_base": unit_base,
            "unit_base_label": item.get_unidade_interna_display() or unit_base or "-",
            "balance_base": balance_base,
            "balance_display": cls._format_quantity_display(balance_base, item, unit_base),
            "consumption_window_base": consumption_window_base,
            "consumption_window_display": cls._format_quantity_display(consumption_window_base, item, unit_base),
            "consumption_average_base": consumption_average_base,
            "consumption_average_display": cls._format_daily_quantity_display(consumption_average_base, item, unit_base),
            "days_remaining": round(days_remaining, 2) if days_remaining is not None and isfinite(days_remaining) else None,
            "days_remaining_display": cls._format_days_remaining(days_remaining),
            "coverage_days": coverage_days,
            "suggested_quantity_base": suggested_quantity_base,
            "suggested_quantity_base_display": cls._format_quantity_display(suggested_quantity_base, item, unit_base),
            "suggested_quantity_input": suggested_quantity_input,
            "suggested_quantity_display": cls._format_request_quantity_display(suggested_quantity_input, request_quantity_contract),
            "estimated_suggestion_value": estimated_suggestion_value,
            "supplier": procurement.get("supplier") or {},
            "price_unit_base": price_unit_base,
            "price_unit_request": request_unit_price,
            "price_display": cls._format_currency_per_request(request_unit_price, request_quantity_contract),
            "price_display_base": cls._format_currency_per_base(price_unit_base, unit_base),
            "price_source": (procurement.get("price") or {}).get("source") or "none",
            "price_source_label": cls.PRICE_SOURCE_LABELS.get((procurement.get("price") or {}).get("source") or "none", "Sem preco coerente"),
            "price_reference_document": (procurement.get("price") or {}).get("document_number"),
            "price_reference_type": (procurement.get("price") or {}).get("document_type"),
            "price_reference_date": (procurement.get("price") or {}).get("document_date"),
            "potential_quotes": [],
            "recommended_potential_quote": None,
            "selected_potential_quote": None,
            "selected_potential_quote_id": "",
            "status": status,
            "status_label": cls.STATUS_LABELS.get(status, status),
            "status_badge_class": cls.STATUS_BADGE_CLASSES.get(status, "secondary"),
            "validation_notes": validation_notes,
            "selected": False,
            "manual_quantity_input": None,
            "manual_quantity_value": cls._format_number_input_value(suggested_quantity_input),
            "request_quantity_mode": request_quantity_contract.get("mode"),
            "request_quantity_factor_base": request_factor_base,
            "request_quantity_unit_label": request_quantity_contract.get("unit_label"),
            "request_quantity_unit_label_plural": request_quantity_contract.get("unit_label_plural"),
            "request_quantity_short_label": request_quantity_contract.get("short_label"),
            "request_quantity_note": request_quantity_contract.get("note"),
            "requested_quantity_base": None,
            "requested_quantity_base_display": None,
            "requested_quantity_input": None,
            "requested_quantity_display": None,
            "requested_total_value": None,
            "selection_note": None,
            "last_saida_at": window_data.get("last_saida_at") or history_data.get("last_saida_at"),
            "procurement_note": procurement.get("note"),
        }

    @classmethod
    def _resolve_procurement(
        cls,
        item: Item,
        *,
        preference: FinanceSupplierPreference | None,
        document_candidate: dict[str, Any] | None,
        finance_candidate: dict[str, Any] | None,
        supplier_lookup: dict[int, FinanceSupplier],
    ) -> dict[str, Any]:
        preferred_supplier = cls._serialize_preferred_supplier(preference, supplier_lookup)
        if preferred_supplier and cls._candidate_matches_supplier(document_candidate, preferred_supplier):
            return {
                "supplier": preferred_supplier,
                "price": {**dict(document_candidate or {}), "source": "document_pair"},
                "note": "Fornecedor preferencial confirmado por documento fiscal.",
            }
        if preferred_supplier and cls._candidate_matches_supplier(finance_candidate, preferred_supplier):
            return {
                "supplier": preferred_supplier,
                "price": {**dict(finance_candidate or {}), "source": "finance_pair"},
                "note": "Fornecedor preferencial confirmado pelo financeiro.",
            }
        if cls._candidate_has_supplier(document_candidate):
            return {
                "supplier": dict((document_candidate or {}).get("supplier") or {}),
                "price": {**dict(document_candidate or {}), "source": "document_pair"},
                "note": "Fornecedor e preco derivados do documento fiscal mais recente coerente.",
            }
        if cls._candidate_has_supplier(finance_candidate):
            return {
                "supplier": dict((finance_candidate or {}).get("supplier") or {}),
                "price": {**dict(finance_candidate or {}), "source": "finance_pair"},
                "note": "Fornecedor e preco derivados do lancamento financeiro coerente mais recente.",
            }
        generic_price = cls._resolve_item_master_price(item)
        if preferred_supplier and generic_price:
            return {
                "supplier": preferred_supplier,
                "price": {**generic_price, "source": "preferred_item_master"},
                "note": "Fornecedor preferencial mantido com preco mestre do item, sem misturar documento de outra origem.",
            }
        if document_candidate and cls._is_positive_number((document_candidate.get("unit_price_base") if isinstance(document_candidate, dict) else None)):
            return {
                "supplier": {},
                "price": {**dict(document_candidate), "source": "document_generic"},
                "note": "Preco documental encontrado, mas sem fornecedor coerente para consolidacao.",
            }
        if generic_price:
            return {
                "supplier": preferred_supplier or {},
                "price": {**generic_price, "source": "item_master"},
                "note": "Preco do cadastro mestre usado por falta de par coerente fornecedor+documento.",
            }
        return {
            "supplier": preferred_supplier or {},
            "price": {"source": "none"},
            "note": "Sem fornecedor e preco coerentes suficientes para consolidacao automatica.",
        }

    @classmethod
    def _attach_potential_quotes(
        cls,
        rows: list[dict[str, Any]],
        potential_quotes: dict[str, list[dict[str, Any]]],
        *,
        potential_quote_ids: dict[str, object] | None = None,
    ) -> None:
        quote_map = {
            str(code or "").strip(): str(value or "").strip()
            for code, value in dict(potential_quote_ids or {}).items()
            if str(code or "").strip()
        }
        for row in rows:
            codigo = str(row.get("codigo_item") or "").strip()
            quotes = list(potential_quotes.get(codigo) or [])
            selected_quote_id = quote_map.get(codigo) or ""
            selected_quote = None
            if selected_quote_id:
                selected_quote = next((quote for quote in quotes if str(quote.get("id") or "") == selected_quote_id), None)
                if selected_quote is None:
                    selected_quote = potential_supplier_service.get_quote(
                        selected_quote_id,
                        reference_price_base=row.get("price_unit_base"),
                    )
                    if selected_quote:
                        quotes.append(selected_quote)
            recommended_quote = quotes[0] if quotes else None
            row["potential_quotes"] = quotes
            row["recommended_potential_quote"] = recommended_quote
            row["selected_potential_quote"] = selected_quote
            row["selected_potential_quote_id"] = str((selected_quote or {}).get("id") or "")

    @classmethod
    def _apply_selection(
        cls,
        rows: list[dict[str, Any]],
        *,
        selected_codes: list[str] | tuple[str, ...] | set[str] | None,
        manual_quantities: dict[str, object] | None,
        potential_quote_ids: dict[str, object] | None = None,
    ) -> dict[str, Any]:
        selected_set = {str(code or "").strip() for code in (selected_codes or []) if str(code or "").strip()}
        manual_map = {str(code or "").strip(): value for code, value in dict(manual_quantities or {}).items() if str(code or "").strip()}
        potential_map = {str(code or "").strip(): str(value or "").strip() for code, value in dict(potential_quote_ids or {}).items() if str(code or "").strip()}
        groups: dict[str, dict[str, Any]] = {}
        category_groups: dict[str, dict[str, Any]] = {}
        messages: list[str] = []
        selected_count = 0
        total_value = 0.0
        total_quantity = 0.0

        for row in rows:
            codigo = str(row.get("codigo_item") or "").strip()
            if codigo not in selected_set:
                continue
            row["selected"] = True
            manual_raw = manual_map.get(codigo)
            manual_value, manual_error = cls._parse_manual_quantity(manual_raw)
            if manual_error:
                row["selection_note"] = cls.VALIDATION_MESSAGES[manual_error]
                row["validation_notes"].append(
                    {
                        "code": manual_error,
                        "message": cls.VALIDATION_MESSAGES[manual_error],
                        "blocking": False,
                    }
                )
                messages.append(f"{codigo}: {cls.VALIDATION_MESSAGES[manual_error]}")
            requested_quantity_input = manual_value if manual_value is not None else float(row.get("suggested_quantity_input") or 0.0)
            request_factor_base = float(row.get("request_quantity_factor_base") or 1.0)
            requested_quantity = requested_quantity_input * request_factor_base
            row["manual_quantity_input"] = manual_raw
            row["manual_quantity_value"] = cls._format_number_input_value(requested_quantity_input)
            row["requested_quantity_input"] = requested_quantity_input
            if requested_quantity <= cls.BALANCE_TOLERANCE:
                row["requested_quantity_base"] = 0.0
                row["requested_quantity_base_display"] = cls._format_quantity_display(0.0, None, row.get("unit_base"))
                row["requested_quantity_display"] = cls._format_request_quantity_display(0.0, row)
                if manual_raw not in (None, ""):
                    messages.append(f"{codigo}: {cls.VALIDATION_MESSAGES['quantidade_manual_zerada']}")
                continue

            selected_count += 1
            row["requested_quantity_base"] = requested_quantity
            row["requested_quantity_base_display"] = cls._format_quantity_display(requested_quantity, None, row.get("unit_base"))
            row["requested_quantity_display"] = cls._format_request_quantity_display(requested_quantity_input, row)
            price_unit_base = row.get("price_unit_base")
            if cls._is_positive_number(price_unit_base):
                row["requested_total_value"] = round(float(price_unit_base) * requested_quantity, 2)
                total_value += float(row["requested_total_value"] or 0.0)
            else:
                row["requested_total_value"] = None
            selected_quote_id = potential_map.get(codigo) or str(row.get("selected_potential_quote_id") or "").strip()
            selected_quote = None
            if selected_quote_id:
                selected_quote = next(
                    (
                        quote
                        for quote in list(row.get("potential_quotes") or [])
                        if str(quote.get("id") or "") == selected_quote_id
                    ),
                    None,
                )
            if selected_quote:
                quote_payload = dict(selected_quote)
                quote_price_base = quote_payload.get("unit_price_base")
                if cls._is_positive_number(quote_price_base):
                    potential_total = round(float(quote_price_base) * requested_quantity, 2)
                    quote_payload["requested_total_value"] = potential_total
                    quote_payload["requested_total_display"] = cls._format_brl(potential_total)
                    if cls._is_positive_number(price_unit_base):
                        savings_total = round((float(price_unit_base) - float(quote_price_base)) * requested_quantity, 2)
                        quote_payload["estimated_savings_total"] = savings_total
                        quote_payload["estimated_savings_display"] = cls._format_brl(savings_total)
                row["selected_potential_quote"] = quote_payload
                row["selected_potential_quote_id"] = str(quote_payload.get("id") or "")
            else:
                row["selected_potential_quote"] = None
                row["selected_potential_quote_id"] = ""
            total_quantity += requested_quantity

            supplier = dict(row.get("supplier") or {})
            supplier_key = cls._supplier_group_key(supplier)
            groups.setdefault(
                supplier_key,
                {
                    "supplier_key": supplier_key,
                    "supplier": supplier,
                    "supplier_name": supplier.get("name") or "Sem fornecedor identificado",
                    "items": [],
                    "requested_quantity_total": 0.0,
                    "requested_value_total": 0.0,
                },
            )
            groups[supplier_key]["items"].append(row)
            groups[supplier_key]["requested_quantity_total"] += requested_quantity
            groups[supplier_key]["requested_value_total"] += float(row.get("requested_total_value") or 0.0)

            category_key = cls._category_group_key(row.get("categoria"))
            category_groups.setdefault(
                category_key,
                {
                    "category_key": category_key,
                    "category_name": cls._category_display_name(row.get("categoria")),
                    "items": [],
                    "requested_quantity_total": 0.0,
                    "requested_value_total": 0.0,
                },
            )
            category_groups[category_key]["items"].append(row)
            category_groups[category_key]["requested_quantity_total"] += requested_quantity
            category_groups[category_key]["requested_value_total"] += float(row.get("requested_total_value") or 0.0)

        groups_list = sorted(
            groups.values(),
            key=lambda group: (
                1 if str(group.get("supplier_name") or "").strip().lower() == "sem fornecedor identificado" else 0,
                str(group.get("supplier_name") or "").lower(),
            ),
        )
        for group in groups_list:
            group["items"] = sorted(
                list(group.get("items") or []),
                key=lambda item: (
                    cls._category_group_key(item.get("categoria")),
                    str(item.get("descricao") or "").casefold(),
                    str(item.get("codigo_item") or "").casefold(),
                ),
            )

        category_groups_list = sorted(
            category_groups.values(),
            key=lambda group: cls._category_group_key(group.get("category_name")),
        )
        for group in category_groups_list:
            group["items"] = sorted(
                list(group.get("items") or []),
                key=lambda item: (
                    str((item.get("supplier") or {}).get("name") or "Sem fornecedor identificado").casefold(),
                    str(item.get("descricao") or "").casefold(),
                    str(item.get("codigo_item") or "").casefold(),
                ),
            )

        return {
            "selected_count": selected_count,
            "group_count": len(groups_list),
            "category_group_count": len(category_groups_list),
            "requested_quantity_total": total_quantity,
            "requested_value_total": round(total_value, 2),
            "groups": groups_list,
            "category_groups": category_groups_list,
            "messages": messages,
        }

    @classmethod
    def _build_summary(
        cls,
        rows: list[dict[str, Any]],
        visible_rows: list[dict[str, Any]],
        cart: dict[str, Any],
    ) -> dict[str, Any]:
        return {
            "total_items": len(rows),
            "visible_items": len(visible_rows),
            "critical_items": sum(1 for row in visible_rows if row.get("status") == cls.STATUS_CRITICAL),
            "attention_items": sum(1 for row in visible_rows if row.get("status") == cls.STATUS_ATTENTION),
            "no_history_items": sum(1 for row in visible_rows if row.get("status") == cls.STATUS_NO_HISTORY),
            "selected_items": int(cart.get("selected_count") or 0),
            "selected_groups": int(cart.get("group_count") or 0),
            "selected_categories": int(cart.get("category_group_count") or 0),
            "selected_quantity_total": float(cart.get("requested_quantity_total") or 0.0),
            "selected_value_total": float(cart.get("requested_value_total") or 0.0),
        }

    @classmethod
    def _build_visible_category_groups(cls, rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
        groups: dict[str, dict[str, Any]] = {}
        for row in rows:
            category_key = cls._category_group_key(row.get("categoria"))
            groups.setdefault(
                category_key,
                {
                    "category_key": category_key,
                    "category_name": cls._category_display_name(row.get("categoria")),
                    "visible_count": 0,
                    "selected_count": 0,
                    "requested_quantity_total": 0.0,
                    "requested_value_total": 0.0,
                },
            )
            groups[category_key]["visible_count"] += 1
            if row.get("selected"):
                groups[category_key]["selected_count"] += 1
                groups[category_key]["requested_quantity_total"] += float(row.get("requested_quantity_base") or 0.0)
                groups[category_key]["requested_value_total"] += float(row.get("requested_total_value") or 0.0)
        return sorted(groups.values(), key=lambda group: str(group.get("category_key") or ""))

    @staticmethod
    def _category_display_name(value: object) -> str:
        return str(value or "Sem categoria").strip() or "Sem categoria"

    @classmethod
    def _category_group_key(cls, value: object) -> str:
        return cls._category_display_name(value).casefold()

    @classmethod
    def _classify_row_status(
        cls,
        *,
        balance_base: float,
        consumption_average_base: float,
        suggested_quantity_base: float,
        validation_codes: list[str],
    ) -> str:
        if any(code in cls.BLOCKING_VALIDATIONS for code in validation_codes):
            return cls.STATUS_CRITICAL
        if balance_base <= cls.BALANCE_TOLERANCE and consumption_average_base > cls.BALANCE_TOLERANCE:
            return cls.STATUS_CRITICAL
        if "item_sem_historico" in validation_codes:
            return cls.STATUS_NO_HISTORY
        if "item_sem_consumo" in validation_codes:
            return cls.STATUS_NO_CONSUMPTION
        if suggested_quantity_base > cls.BALANCE_TOLERANCE or "item_sem_fornecedor" in validation_codes:
            return cls.STATUS_ATTENTION
        return cls.STATUS_OK

    @classmethod
    def _status_matches_filter(cls, status_filter: str, row: dict[str, Any]) -> bool:
        if status_filter == cls.STATUS_ALL:
            return True
        if status_filter == cls.STATUS_ACTIONABLE:
            return str(row.get("status") or "") in {cls.STATUS_CRITICAL, cls.STATUS_ATTENTION}
        return str(row.get("status") or "") == status_filter

    @classmethod
    def _row_sort_key(cls, row: dict[str, Any]) -> tuple[Any, ...]:
        priority = {
            cls.STATUS_CRITICAL: 0,
            cls.STATUS_ATTENTION: 1,
            cls.STATUS_NO_HISTORY: 2,
            cls.STATUS_NO_CONSUMPTION: 3,
            cls.STATUS_OK: 4,
        }
        days_remaining = row.get("days_remaining")
        normalized_days = float(days_remaining) if days_remaining is not None else 999999.0
        return (
            priority.get(str(row.get("status") or ""), 9),
            normalized_days,
            -float(row.get("suggested_quantity_base") or 0.0),
            str(row.get("descricao") or "").lower(),
        )

    @classmethod
    def _resolve_item_master_price(cls, item: Item) -> dict[str, Any] | None:
        if cls._is_positive_number(getattr(item, "preco_compra_unitario_base", None)):
            return {
                "unit_price_base": float(item.preco_compra_unitario_base or 0.0),
                "document_number": getattr(item, "preco_compra_documento", None),
                "document_type": "item_compra",
                "document_date": item.preco_compra_data_recebimento.isoformat() if getattr(item, "preco_compra_data_recebimento", None) else (item.preco_compra_data_emissao.isoformat() if getattr(item, "preco_compra_data_emissao", None) else None),
            }
        if cls._is_positive_number(getattr(item, "preco_reposicao_unitario_base", None)):
            return {
                "unit_price_base": float(item.preco_reposicao_unitario_base or 0.0),
                "document_number": getattr(item, "preco_reposicao_query", None),
                "document_type": "item_reposicao",
                "document_date": None,
            }
        return None

    @classmethod
    def _serialize_preferred_supplier(
        cls,
        preference: FinanceSupplierPreference | None,
        supplier_lookup: dict[int, FinanceSupplier],
    ) -> dict[str, Any]:
        if preference is None or not preference.fornecedor_id:
            return {}
        supplier = supplier_lookup.get(int(preference.fornecedor_id))
        if supplier is None:
            return {
                "id": int(preference.fornecedor_id),
                "name": None,
                "cnpj": None,
            }
        return {
            "id": int(supplier.id),
            "name": supplier.nome_exibicao(),
            "cnpj": supplier.cnpj,
        }

    @classmethod
    def _candidate_has_supplier(cls, candidate: dict[str, Any] | None) -> bool:
        supplier = (candidate or {}).get("supplier") if isinstance(candidate, dict) else {}
        return bool((supplier or {}).get("id") or (supplier or {}).get("name"))

    @classmethod
    def _candidate_matches_supplier(cls, candidate: dict[str, Any] | None, supplier: dict[str, Any] | None) -> bool:
        candidate_supplier = (candidate or {}).get("supplier") if isinstance(candidate, dict) else {}
        supplier = supplier or {}
        supplier_id = supplier.get("id")
        candidate_id = (candidate_supplier or {}).get("id")
        if supplier_id and candidate_id:
            return int(supplier_id) == int(candidate_id)
        supplier_name = str(supplier.get("name") or "").strip().lower()
        candidate_name = str((candidate_supplier or {}).get("name") or "").strip().lower()
        return bool(supplier_name and candidate_name and supplier_name == candidate_name)

    @classmethod
    def _parse_manual_quantity(cls, value: object) -> tuple[float | None, str | None]:
        if value in (None, ""):
            return None, None
        text = str(value).strip().replace(",", ".")
        try:
            parsed = float(text)
        except (TypeError, ValueError):
            return None, "quantidade_manual_invalida"
        if not isfinite(parsed) or parsed < 0:
            return None, "quantidade_manual_invalida"
        return parsed, None

    @classmethod
    def _build_request_quantity_contract(cls, item: Item, unit_base: object) -> dict[str, Any]:
        base_unit = cls._normalize_unit(unit_base)
        base_label = str(item.get_unidade_interna_display() or base_unit or "base").strip() or "base"
        package_type = str(getattr(item, "tipo_embalagem_novo", None) or getattr(item, "tipo_embalagem", None) or "").strip().lower()
        unit_raw = cls._normalize_unit(getattr(item, "unidade", None))
        inferred_measure = infer_packaging_measure(item)
        inferred_factor = float((inferred_measure or (0.0, None))[0] or 0.0) if inferred_measure else 0.0
        factor_base = float(resolve_packaging_factor(item) or 0.0)
        if inferred_factor > cls.BALANCE_TOLERANCE:
            factor_base = inferred_factor

        use_request_units = False
        helper_context: str | None = None
        if base_unit == "un":
            use_request_units = True
            factor_base = 1.0
        elif package_type and factor_base > cls.BALANCE_TOLERANCE:
            use_request_units = True
            if package_type not in {"", "litro"}:
                try:
                    helper_context = item.get_nome_embalagem()
                except Exception:
                    helper_context = package_type
        elif unit_raw in cls.REQUEST_UNIT_ALIASES and factor_base > cls.BALANCE_TOLERANCE:
            use_request_units = True

        if not use_request_units or factor_base <= cls.BALANCE_TOLERANCE:
            return {
                "mode": "base",
                "factor_base": 1.0,
                "unit_label": base_label,
                "unit_label_plural": base_label,
                "short_label": base_label,
                "note": f"pedido em unidade base: {base_label}",
            }

        factor_display = cls._format_simple_value(factor_base, base_label)
        note = f"1 unidade = {factor_display}"
        if helper_context:
            note = f"{note} ({helper_context})"
        return {
            "mode": "request_unit",
            "factor_base": factor_base,
            "unit_label": "unidade",
            "unit_label_plural": "unidades",
            "short_label": "un",
            "note": note,
        }

    @classmethod
    def _normalize_int(cls, value: object, *, default: int, minimum: int, maximum: int) -> int:
        try:
            parsed = int(str(value).strip()) if value not in (None, "") else default
        except (TypeError, ValueError):
            return default
        return max(minimum, min(maximum, parsed))

    @classmethod
    def _normalize_bool(cls, value: object) -> bool:
        if isinstance(value, bool):
            return value
        return str(value or "").strip().lower() in {"1", "true", "on", "yes", "sim"}

    @classmethod
    def _normalize_status(cls, value: object) -> str:
        normalized = str(value or cls.STATUS_ALL).strip().lower()
        if normalized not in cls.FILTER_STATUS_OPTIONS:
            return cls.STATUS_ALL
        return normalized

    @classmethod
    def _normalize_unit(cls, value: object) -> str:
        return str(value or "").strip().lower()

    @classmethod
    def _format_quantity_display(cls, value: object, item: Item | None, unit_base: object) -> str:
        unit_hint = cls._normalize_unit(unit_base)
        try:
            parsed = float(value or 0.0)
        except (TypeError, ValueError):
            parsed = 0.0
        if abs(parsed) <= cls.BALANCE_TOLERANCE:
            return cls._format_simple_value(0.0, unit_hint)
        if parsed > 0:
            display = OperationVisualPayloadService.format_balance_display(parsed, item, unit_hint=unit_hint, short=True)
            if display:
                return display
        if parsed < 0:
            display = OperationVisualPayloadService.format_balance_display(abs(parsed), item, unit_hint=unit_hint, short=True)
            if display:
                return f"-{display}"
        return cls._format_simple_value(parsed, unit_hint)

    @classmethod
    def _format_daily_quantity_display(cls, value: object, item: Item | None, unit_base: object) -> str:
        return f"{cls._format_quantity_display(value, item, unit_base)}/dia"

    @classmethod
    def _format_days_remaining(cls, value: object) -> str:
        if value is None:
            return "Sem consumo"
        try:
            parsed = float(value)
        except (TypeError, ValueError):
            return "Sem consumo"
        if not isfinite(parsed):
            return "Sem consumo"
        if abs(parsed - round(parsed)) <= 1e-6:
            return f"{int(round(parsed))} dias"
        return f"{parsed:.2f} dias"

    @classmethod
    def _format_currency_per_base(cls, value: object, unit_base: object) -> str | None:
        if not cls._is_positive_number(value):
            return None
        amount = float(value or 0.0)
        formatted = f"R$ {amount:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        return f"{formatted}/{cls._normalize_unit(unit_base) or 'base'}"

    @classmethod
    def _format_brl(cls, value: object) -> str:
        try:
            amount = float(value or 0.0)
        except (TypeError, ValueError):
            amount = 0.0
        return f"R$ {amount:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

    @classmethod
    def _format_currency_per_request(cls, value: object, request_contract: dict[str, Any] | None) -> str | None:
        if not cls._is_positive_number(value):
            return None
        amount = float(value or 0.0)
        formatted = f"R$ {amount:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        label = str((request_contract or {}).get("unit_label") or "unidade").strip() or "unidade"
        return f"{formatted}/{label}"

    @classmethod
    def _format_simple_value(cls, value: float, unit_hint: str | None) -> str:
        parsed = float(value or 0.0)
        if abs(parsed - round(parsed)) <= 1e-6:
            number = str(int(round(parsed)))
        else:
            number = f"{parsed:.6f}".rstrip("0").rstrip(".")
        unit = str(unit_hint or "").strip()
        return f"{number} {unit}".strip()

    @classmethod
    def _format_request_quantity_display(cls, value: object, request_contract: dict[str, Any] | None) -> str:
        try:
            parsed = float(value or 0.0)
        except (TypeError, ValueError):
            parsed = 0.0
        if abs(parsed) <= cls.BALANCE_TOLERANCE:
            parsed = 0.0
        number = cls._format_number_input_value(parsed)
        singular = str(
            (request_contract or {}).get("unit_label")
            or (request_contract or {}).get("request_quantity_unit_label")
            or "unidade"
        ).strip() or "unidade"
        plural = str(
            (request_contract or {}).get("unit_label_plural")
            or (request_contract or {}).get("request_quantity_unit_label_plural")
            or singular
        ).strip() or singular
        label = singular if abs(parsed - 1.0) <= 1e-6 else plural
        return f"{number} {label}".strip()

    @classmethod
    def _format_number_input_value(cls, value: object) -> str:
        try:
            parsed = float(value or 0.0)
        except (TypeError, ValueError):
            parsed = 0.0
        if abs(parsed - round(parsed)) <= 1e-6:
            return str(int(round(parsed)))
        return f"{parsed:.6f}".rstrip("0").rstrip(".")

    @classmethod
    def _style_sheet_header(cls, row) -> None:
        fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        font = Font(color="FFFFFF", bold=True)
        border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )
        for cell in row:
            cell.fill = fill
            cell.font = font
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    @classmethod
    def _append_group_rows(cls, sheet, group: dict[str, Any], filters: dict[str, Any]) -> None:
        border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )
        supplier_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
        total_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
        supplier_name = str(group.get("supplier_name") or "Sem fornecedor identificado")

        title_row = sheet.max_row + 1
        sheet.append([supplier_name])
        sheet.merge_cells(start_row=title_row, start_column=1, end_row=title_row, end_column=18)
        title_cell = sheet.cell(row=title_row, column=1)
        title_cell.fill = supplier_fill
        title_cell.font = Font(bold=True, color="1E3A8A")
        title_cell.border = border
        title_cell.alignment = Alignment(horizontal="left", vertical="center")

        for item in group.get("items") or []:
            validations = "; ".join(note.get("message") or note.get("code") for note in item.get("validation_notes") or [])
            sheet.append(
                [
                    supplier_name,
                    ((group.get("supplier") or {}).get("cnpj") or ""),
                    item.get("codigo_item") or "",
                    item.get("descricao") or "",
                    item.get("categoria") or "",
                    item.get("unit_base") or "",
                    float(item.get("balance_base") or 0.0),
                    float(item.get("consumption_window_base") or 0.0),
                    float(item.get("consumption_average_base") or 0.0),
                    item.get("days_remaining"),
                    float(filters.get("coverage_days") or cls.DEFAULT_COVERAGE_DAYS),
                    float(item.get("suggested_quantity_base") or 0.0),
                    float(item.get("requested_quantity_base") or 0.0),
                    item.get("price_unit_base"),
                    item.get("requested_total_value"),
                    item.get("price_source_label") or "",
                    item.get("status_label") or "",
                    validations,
                ]
            )
            current_row = sheet.max_row
            for col_idx in range(1, 19):
                cell = sheet.cell(row=current_row, column=col_idx)
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if col_idx in {14, 15} and isinstance(cell.value, (int, float)):
                    cell.number_format = 'R$ #,##0.00'

        total_row = sheet.max_row + 1
        sheet.append(
            [
                f"Total fornecedor: {supplier_name}",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                float(group.get("requested_quantity_total") or 0.0),
                "",
                float(group.get("requested_value_total") or 0.0),
                "",
                "",
                "",
            ]
        )
        for col_idx in range(1, 19):
            cell = sheet.cell(row=total_row, column=col_idx)
            cell.fill = total_fill
            cell.font = Font(bold=True)
            cell.border = border
            if col_idx in {14, 15} and isinstance(cell.value, (int, float)):
                cell.number_format = 'R$ #,##0.00'
        sheet.append([])

    @classmethod
    def _fit_sheet_columns(cls, sheet) -> None:
        widths = {
            1: 28,
            2: 18,
            3: 16,
            4: 42,
            5: 20,
            6: 14,
            7: 14,
            8: 16,
            9: 18,
            10: 15,
            11: 15,
            12: 16,
            13: 16,
            14: 18,
            15: 18,
            16: 24,
            17: 14,
            18: 44,
        }
        for col_idx, width in widths.items():
            sheet.column_dimensions[get_column_letter(col_idx)].width = width

    @classmethod
    def _supplier_group_key(cls, supplier: dict[str, Any]) -> str:
        supplier_id = supplier.get("id")
        if supplier_id:
            return f"supplier:{int(supplier_id)}"
        supplier_name = str(supplier.get("name") or "").strip()
        if supplier_name:
            return f"name:{supplier_name.lower()}"
        return "supplier:none"

    @staticmethod
    def _is_positive_number(value: object) -> bool:
        try:
            parsed = float(value)
        except (TypeError, ValueError):
            return False
        return isfinite(parsed) and parsed > 0


purchase_projection_service = PurchaseProjectionService()