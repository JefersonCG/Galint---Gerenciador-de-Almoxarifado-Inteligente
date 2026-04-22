"""Serviço de geração de Relatório de Escopo - análise executiva multinível."""
from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func

from ..extensions import db
from ..models import DocumentoEntradaEstoqueItem, Entrada, FinanceLedgerEntry, InventarioEvento, Item, Saida, StockBalance, StockMovement
from ..utils.report_branding import get_company_header_lines
from .legacy_stock_normalizer import ignore_packaging_metadata_for_stock, resolve_canonical_unit, resolve_packaging_factor
from .price_normalization import infer_document_quantity_unit_for_item, normalize_document_line, should_autofix_packaged_document_unit


class ScopeReportService:
    """Geração de relatório executivo de escopo com visão financeira e timeline."""

    COLOR_HEADER = "0066CC"
    COLOR_HEADER_DARK = "1E3A8A"
    COLOR_SUBHEADER = "0F62FE"
    COLOR_LIGHT = "F6F8FC"
    COLOR_LIGHT_ALT = "EEF4FF"
    COLOR_PANEL = "E8F1FF"
    COLOR_WHITE = "FFFFFF"
    COLOR_TEXT = "0F172A"
    COLOR_MUTED = "475569"
    COLOR_SUCCESS = "15803D"
    COLOR_WARNING = "B45309"
    COLOR_DANGER = "B91C1C"
    COLOR_ALERT = "FEF3C7"
    ANALYSIS_TOLERANCE = 1e-6
    BORDER_THIN = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    @classmethod
    def generate_executive_scope_report(cls, category_name: str | None = None) -> BytesIO:
        """Gera relatório executivo completo de escopo."""
        scope_category = cls._normalize_scope_category(category_name)
        scope_label = scope_category or "Geral"
        scoped_items = cls._load_scope_items(scope_category)
        if not scoped_items:
            if scope_category:
                raise ValueError(f"Nenhum item encontrado para a categoria '{scope_category}'.")
            raise ValueError("Nenhum item encontrado para gerar o relatório.")

        items = cls._collect_inventory_data(scoped_items)
        category_analysis = cls._analyze_by_category(items)
        brand_analysis = cls._analyze_by_brand(items)
        rankings = cls._generate_rankings(items)
        inconsistency_analysis = cls._build_inconsistency_analysis(scoped_items, items)
        timeline = cls._build_timeline(scoped_items)

        wb = Workbook()
        wb.remove(wb.active)
        cls._create_executive_summary_sheet(wb, category_analysis, items, timeline, scope_label)
        cls._create_timeline_sheet(wb, timeline, scope_label)
        cls._create_brand_analysis_sheet(wb, brand_analysis, scope_label)
        cls._create_rankings_sheet(wb, rankings, scope_label)
        cls._create_inconsistency_sheet(wb, inconsistency_analysis, scope_label)
        cls._create_detailed_sheet(wb, items, scope_label)

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    @classmethod
    def _normalize_scope_category(cls, category_name: str | None) -> str | None:
        normalized = (category_name or "").strip()
        return normalized or None

    @classmethod
    def _load_scope_items(cls, category_name: str | None = None) -> list[Item]:
        scope_category = cls._normalize_scope_category(category_name)
        scoped_items: list[Item] = []
        for item_obj in Item.query.order_by(Item.categoria.asc(), Item.descricao.asc()).all():
            category_label = str(item_obj.categoria or "Sem categoria").strip() or "Sem categoria"
            if scope_category and category_label.casefold() != scope_category.casefold():
                continue
            scoped_items.append(item_obj)
        return scoped_items

    @staticmethod
    def _safe_float(value: object) -> float:
        try:
            return float(value or 0.0)
        except (TypeError, ValueError):
            return 0.0

    @staticmethod
    def _format_datetime(value: object) -> str:
        if isinstance(value, datetime):
            return value.strftime("%d/%m/%Y %H:%M")
        return "-"

    @classmethod
    def _normalize_historical_price_candidate(
        cls,
        item: Item,
        *,
        quantity: object,
        quantity_base: object,
        quantity_unit: object,
        price_unit: object,
        unit_price: object,
        total_price: object,
    ) -> dict[str, Any] | None:
        quantity_value = cls._safe_float(quantity)
        if quantity_value <= 0:
            return None

        stored_quantity_unit = (str(quantity_unit or "").strip().lower())
        effective_quantity_unit = stored_quantity_unit
        auto_fixed = False

        if not effective_quantity_unit:
            effective_quantity_unit = infer_document_quantity_unit_for_item(item)
            auto_fixed = True
        elif should_autofix_packaged_document_unit(
            item,
            current_unit=effective_quantity_unit,
            quantity=quantity_value,
            quantity_base=quantity_base,
        ):
            effective_quantity_unit = infer_document_quantity_unit_for_item(item)
            auto_fixed = True

        stored_price_unit = (str(price_unit or "").strip().lower())
        if auto_fixed and (not stored_price_unit or stored_price_unit == stored_quantity_unit):
            effective_price_unit = effective_quantity_unit
        else:
            effective_price_unit = stored_price_unit or effective_quantity_unit

        try:
            normalized = normalize_document_line(
                item,
                quantity=quantity_value,
                quantity_unit=effective_quantity_unit,
                unit_price=float(unit_price) if unit_price not in (None, "") else None,
                total_price=float(total_price) if total_price not in (None, "") else None,
                price_unit=effective_price_unit,
            )
        except Exception:
            return None

        unit_price_base = cls._safe_float(normalized.unit_price_base)
        if unit_price_base <= 0:
            return None

        return {
            "unit_price_base": unit_price_base,
            "price_unit": normalized.price_unit,
            "factor_to_base": cls._safe_float(normalized.factor_to_base) or 1.0,
            "auto_fixed": auto_fixed,
        }

    @classmethod
    def _build_price_override_map(cls, item_map: dict[str, Item]) -> dict[str, dict[str, Any]]:
        if not item_map:
            return {}

        codes = sorted(item_map)
        overrides: dict[str, dict[str, Any]] = {}

        def consider(code: str, candidate: dict[str, Any]) -> None:
            current = overrides.get(code)
            rank = (int(candidate.get("priority", 0)), int(candidate.get("sort_key", 0)))
            if current is None or rank > current["_rank"]:
                candidate["_rank"] = rank
                overrides[code] = candidate

        for row in (
            DocumentoEntradaEstoqueItem.query
            .filter(DocumentoEntradaEstoqueItem.codigo_item.in_(codes))
            .order_by(DocumentoEntradaEstoqueItem.id_documento_item.asc())
            .all()
        ):
            code = str(row.codigo_item or "").strip()
            item = item_map.get(code)
            if item is None:
                continue
            normalized = cls._normalize_historical_price_candidate(
                item,
                quantity=row.quantidade,
                quantity_base=row.quantidade_base,
                quantity_unit=row.unidade_quantidade,
                price_unit=row.unidade_preco,
                unit_price=row.valor_unitario,
                total_price=row.valor_total,
            )
            if not normalized:
                continue
            auto_fixed = bool(normalized["auto_fixed"])
            consider(code, {
                "kind": "real",
                "origin": "NF/Cupom (autocorreção de embalagem)" if auto_fixed else "NF/Cupom",
                "unit_price_base": normalized["unit_price_base"],
                "price_unit": normalized["price_unit"],
                "factor_to_base": normalized["factor_to_base"],
                "warning": "Preço recalculado a partir do histórico documental de embalagem." if auto_fixed else None,
                "priority": 40 if auto_fixed else 20,
                "sort_key": int(row.id_documento_item or 0),
            })

        for row in (
            FinanceLedgerEntry.query
            .filter(FinanceLedgerEntry.codigo_item.in_(codes))
            .order_by(FinanceLedgerEntry.id.asc())
            .all()
        ):
            code = str(row.codigo_item or "").strip()
            item = item_map.get(code)
            if item is None:
                continue
            normalized = cls._normalize_historical_price_candidate(
                item,
                quantity=row.quantidade,
                quantity_base=row.quantidade_base,
                quantity_unit=row.unidade_quantidade,
                price_unit=row.unidade_preco,
                unit_price=row.valor_unitario,
                total_price=row.valor_total,
            )
            if not normalized:
                continue
            auto_fixed = bool(normalized["auto_fixed"])
            consider(code, {
                "kind": "real",
                "origin": "Ledger financeiro (autocorreção de embalagem)" if auto_fixed else "Ledger financeiro",
                "unit_price_base": normalized["unit_price_base"],
                "price_unit": normalized["price_unit"],
                "factor_to_base": normalized["factor_to_base"],
                "warning": "Preço recalculado a partir do histórico financeiro de embalagem." if auto_fixed else None,
                "priority": 30 if auto_fixed else 10,
                "sort_key": int(row.id or 0),
            })

        return {code: {k: v for k, v in data.items() if k != "_rank"} for code, data in overrides.items()}

    @classmethod
    def _is_packaging_price_suspect(cls, item: Item, *, kind: str) -> bool:
        raw_attr = f"preco_{kind}_unitario"
        base_attr = f"preco_{kind}_unitario_base"
        unit_attr = f"preco_{kind}_unidade_preco"
        factor_attr = f"preco_{kind}_fator_base"

        raw_value = cls._safe_float(getattr(item, raw_attr, None))
        base_value = cls._safe_float(getattr(item, base_attr, None))
        stored_factor = cls._safe_float(getattr(item, factor_attr, None))
        packaging_factor = cls._safe_float(resolve_packaging_factor(item))
        stored_unit = (str(getattr(item, unit_attr, None) or "").strip().lower())
        candidate_units = {
            (str(item.unidade or "").strip().lower()),
            (str(resolve_canonical_unit(item) or "").strip().lower()),
        }
        candidate_units.discard("")

        if raw_value <= 0 or base_value <= 0:
            return False
        if packaging_factor <= 1 or ignore_packaging_metadata_for_stock(item):
            return False
        if abs(base_value - raw_value) > 1e-8:
            return False
        if stored_factor not in (0.0, 1.0):
            return False
        if stored_unit not in candidate_units:
            return False
        return True

    @classmethod
    def _resolve_item_price_context(cls, item: Item, override: dict[str, Any] | None) -> dict[str, Any]:
        kind = None
        unit_price_base = None
        origin = "Sem preço"
        price_unit = None
        factor_to_base = None
        warning = None
        override_applied = False

        compra_base = getattr(item, "preco_compra_unitario_base", None) or getattr(item, "preco_compra_unitario", None)
        reposicao_base = getattr(item, "preco_reposicao_unitario_base", None) or getattr(item, "preco_reposicao_unitario", None)

        if cls._safe_float(compra_base) > 0:
            unit_price_base = cls._safe_float(compra_base)
            price_unit = (getattr(item, "preco_compra_unidade_preco", None) or "").strip().lower() or None
            factor_to_base = cls._safe_float(getattr(item, "preco_compra_fator_base", None)) or 1.0
            if getattr(item, "preco_compra_documento", None):
                kind = "real"
                origin = f"NF/Cupom: {item.preco_compra_documento}"
            else:
                kind = "estimado"
                fonte_compra = (getattr(item, "preco_compra_fonte", None) or "cadastro manual").strip()
                origin = f"Estimado ({fonte_compra})"
        elif cls._safe_float(reposicao_base) > 0:
            kind = "estimado"
            unit_price_base = cls._safe_float(reposicao_base)
            price_unit = (getattr(item, "preco_reposicao_unidade_preco", None) or "").strip().lower() or None
            factor_to_base = cls._safe_float(getattr(item, "preco_reposicao_fator_base", None)) or 1.0
            fonte = getattr(item, "preco_reposicao_fonte", None) or "Desconhecida"
            origin = f"Especulativo ({fonte})"

        suspect = False
        if kind == "real":
            suspect = cls._is_packaging_price_suspect(item, kind="compra")
        elif kind == "estimado" and cls._safe_float(compra_base) > 0:
            suspect = cls._is_packaging_price_suspect(item, kind="compra")

        if override and (suspect or unit_price_base is None):
            unit_price_base = cls._safe_float(override.get("unit_price_base"))
            price_unit = override.get("price_unit") or price_unit
            factor_to_base = cls._safe_float(override.get("factor_to_base")) or factor_to_base or 1.0
            kind = override.get("kind") or kind or "real"
            origin = override.get("origin") or origin
            warning = override.get("warning")
            override_applied = True
        elif suspect:
            warning = "Preço base armazenado sugere embalagem ainda não normalizada."

        return {
            "kind": kind,
            "unit_price_base": unit_price_base,
            "origin": origin,
            "price_unit": price_unit,
            "factor_to_base": factor_to_base,
            "warning": warning,
            "override_applied": override_applied,
        }

    @classmethod
    def _collect_inventory_data(cls, scoped_items: list[Item]) -> list[dict[str, Any]]:
        """Coleta dados do inventário com visão física por item e visão monetária saneada."""
        item_map = {str(item.codigo_item): item for item in scoped_items if item.codigo_item}
        price_overrides = cls._build_price_override_map(item_map)

        result: list[dict[str, Any]] = []
        for item_obj in scoped_items:
            code = str(item_obj.codigo_item or "").strip()
            category_label = str(item_obj.categoria or "Sem categoria").strip() or "Sem categoria"
            marca = str(item_obj.marca or "").strip()
            saldo_base = cls._safe_float(item_obj.get_saldo_fisico_total())
            saldo_display = item_obj.get_saldo_fisico_display()
            entrada_eventos = len(item_obj.entradas or [])
            saida_eventos = len(item_obj.saidas or [])
            total_movimentacoes = entrada_eventos + saida_eventos
            ultima_entrada_dt = max((row.data_entrada for row in (item_obj.entradas or []) if getattr(row, "data_entrada", None)), default=None)
            ultima_saida_dt = max((row.data_saida for row in (item_obj.saidas or []) if getattr(row, "data_saida", None)), default=None)
            ultima_movimentacao_dt = max([dt for dt in (ultima_entrada_dt, ultima_saida_dt) if dt is not None], default=None)

            price_context = cls._resolve_item_price_context(item_obj, price_overrides.get(code))
            preco_real = price_context["unit_price_base"] if price_context["kind"] == "real" else None
            preco_estimado = price_context["unit_price_base"] if price_context["kind"] == "estimado" else None
            valor_total_real = round(cls._safe_float(preco_real) * saldo_base, 2) if preco_real else 0.0
            valor_total_estimado = round(cls._safe_float(preco_estimado) * saldo_base, 2) if preco_estimado else 0.0
            valor_total = round(valor_total_real + valor_total_estimado, 2)

            result.append({
                "codigo": code,
                "descricao": item_obj.descricao or "",
                "categoria": category_label,
                "marca": marca,
                "unidade": item_obj.unidade or "un",
                "saldo_base": saldo_base,
                "saldo_display": saldo_display,
                "entrada_eventos": entrada_eventos,
                "saida_eventos": saida_eventos,
                "total_movimentacoes": total_movimentacoes,
                "ultima_entrada_dt": ultima_entrada_dt,
                "ultima_saida_dt": ultima_saida_dt,
                "ultima_movimentacao_dt": ultima_movimentacao_dt,
                "ultima_entrada_label": cls._format_datetime(ultima_entrada_dt),
                "ultima_saida_label": cls._format_datetime(ultima_saida_dt),
                "ultima_movimentacao_label": cls._format_datetime(ultima_movimentacao_dt),
                "preco_real": preco_real,
                "preco_estimado": preco_estimado,
                "preco_tipo_label": "Documentado" if price_context["kind"] == "real" else ("Estimado" if price_context["kind"] == "estimado" else "Sem preço"),
                "preco_unitario_base": cls._safe_float(price_context["unit_price_base"]) or None,
                "preco_unidade": price_context["price_unit"] or (resolve_canonical_unit(item_obj) or item_obj.unidade or "un"),
                "origem_preco": price_context["origin"],
                "price_alert": price_context["warning"],
                "has_price_alert": bool(price_context["warning"]),
                "valor_total_real": valor_total_real,
                "valor_total_estimado": valor_total_estimado,
                "valor_total": valor_total,
                "url_fonte": item_obj.preco_reposicao_url or "",
            })

        return result

    @classmethod
    def _build_timeline(cls, scoped_items: list[Item]) -> list[dict[str, Any]]:
        buckets: dict[str, dict[str, Any]] = {}

        def ensure_bucket(dt: datetime) -> dict[str, Any]:
            key = dt.strftime("%Y-%m")
            bucket = buckets.get(key)
            if bucket is None:
                bucket = {
                    "period_key": key,
                    "period_label": dt.strftime("%m/%Y"),
                    "period_start": datetime(dt.year, dt.month, 1),
                    "entradas": 0,
                    "saidas": 0,
                    "itens": set(),
                    "categorias": set(),
                    "ultimo_evento_dt": None,
                    "ultimo_evento_tipo": None,
                    "ultimo_item": None,
                }
                buckets[key] = bucket
            return bucket

        for item in scoped_items:
            code = str(item.codigo_item or "").strip()
            descricao = str(item.descricao or code)
            categoria = str(item.categoria or "Sem categoria").strip() or "Sem categoria"
            for entrada in item.entradas or []:
                dt = getattr(entrada, "data_entrada", None)
                if not isinstance(dt, datetime):
                    continue
                bucket = ensure_bucket(dt)
                bucket["entradas"] += 1
                bucket["itens"].add(code)
                bucket["categorias"].add(categoria)
                if bucket["ultimo_evento_dt"] is None or dt > bucket["ultimo_evento_dt"]:
                    bucket["ultimo_evento_dt"] = dt
                    bucket["ultimo_evento_tipo"] = "Entrada"
                    bucket["ultimo_item"] = descricao
            for saida in item.saidas or []:
                dt = getattr(saida, "data_saida", None)
                if not isinstance(dt, datetime):
                    continue
                bucket = ensure_bucket(dt)
                bucket["saidas"] += 1
                bucket["itens"].add(code)
                bucket["categorias"].add(categoria)
                if bucket["ultimo_evento_dt"] is None or dt > bucket["ultimo_evento_dt"]:
                    bucket["ultimo_evento_dt"] = dt
                    bucket["ultimo_evento_tipo"] = "Saída"
                    bucket["ultimo_item"] = descricao

        timeline = []
        for bucket in sorted(buckets.values(), key=lambda row: row["period_start"]):
            timeline.append({
                "period_key": bucket["period_key"],
                "period_label": bucket["period_label"],
                "entradas": int(bucket["entradas"] or 0),
                "saidas": int(bucket["saidas"] or 0),
                "movimentacoes": int((bucket["entradas"] or 0) + (bucket["saidas"] or 0)),
                "itens_movimentados": len(bucket["itens"] or set()),
                "categorias_ativas": len(bucket["categorias"] or set()),
                "ultimo_evento_label": (
                    f"{bucket['ultimo_evento_tipo']} em {cls._format_datetime(bucket['ultimo_evento_dt'])} | {bucket['ultimo_item']}"
                    if bucket.get("ultimo_evento_dt") else "-"
                ),
            })
        if len(timeline) > 18:
            timeline = timeline[-18:]
        return timeline

    @classmethod
    def _analyze_by_category(cls, items: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
        categories: dict[str, dict[str, Any]] = {}
        for item in items:
            cat = item["categoria"]
            if cat not in categories:
                categories[cat] = {
                    "total_itens": 0,
                    "itens_com_saldo": 0,
                    "valor_real": 0.0,
                    "valor_estimado": 0.0,
                    "entrada_eventos": 0,
                    "saida_eventos": 0,
                    "itens_com_alerta": 0,
                    "ultima_movimentacao_dt": None,
                }

            current = categories[cat]
            current["total_itens"] += 1
            current["itens_com_saldo"] += 1 if cls._safe_float(item["saldo_base"]) > 0 else 0
            current["valor_real"] += cls._safe_float(item["valor_total_real"])
            current["valor_estimado"] += cls._safe_float(item["valor_total_estimado"])
            current["entrada_eventos"] += int(item["entrada_eventos"] or 0)
            current["saida_eventos"] += int(item["saida_eventos"] or 0)
            current["itens_com_alerta"] += 1 if item["has_price_alert"] else 0
            latest_dt = item.get("ultima_movimentacao_dt")
            if latest_dt and (current["ultima_movimentacao_dt"] is None or latest_dt > current["ultima_movimentacao_dt"]):
                current["ultima_movimentacao_dt"] = latest_dt

        return categories

    @classmethod
    def _analyze_by_brand(cls, items: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
        brands: dict[str, dict[str, Any]] = {}
        for item in items:
            categoria = item["categoria"]
            marca = item["marca"] or "Sem marca"
            key = f"{categoria}||{marca}"
            if key not in brands:
                brands[key] = {
                    "categoria": categoria,
                    "marca": marca,
                    "total_itens": 0,
                    "itens_com_saldo": 0,
                    "valor_total": 0.0,
                    "itens_com_alerta": 0,
                }
            current = brands[key]
            current["total_itens"] += 1
            current["itens_com_saldo"] += 1 if cls._safe_float(item["saldo_base"]) > 0 else 0
            current["valor_total"] += cls._safe_float(item["valor_total"])
            current["itens_com_alerta"] += 1 if item["has_price_alert"] else 0
        return brands

    @classmethod
    def _generate_rankings(cls, items: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
        with_value = [item for item in items if cls._safe_float(item["valor_total"]) > 0]
        with_saida = [item for item in items if int(item["saida_eventos"] or 0) > 0]
        with_movement = [item for item in items if int(item["total_movimentacoes"] or 0) > 0]
        with_alert = [item for item in items if item["has_price_alert"]]
        return {
            "mais_caros": sorted(with_value, key=lambda row: row["valor_total"], reverse=True)[:10],
            "mais_usados": sorted(with_saida, key=lambda row: (int(row["saida_eventos"] or 0), row["valor_total"]), reverse=True)[:10],
            "mais_movimentados": sorted(with_movement, key=lambda row: (int(row["total_movimentacoes"] or 0), row["valor_total"]), reverse=True)[:10],
            "com_alerta": sorted(with_alert, key=lambda row: row["valor_total"], reverse=True)[:10],
        }

    @classmethod
    def _build_inconsistency_analysis(
        cls,
        scoped_items: list[Item],
        items: list[dict[str, Any]],
    ) -> dict[str, list[dict[str, Any]]]:
        codes = sorted({str(item.codigo_item or "").strip() for item in scoped_items if str(item.codigo_item or "").strip()})
        if not codes:
            return {"confirmed": [], "probable": [], "future": []}

        summary_by_code = {str(row.get("codigo") or "").strip(): row for row in items if str(row.get("codigo") or "").strip()}

        entradas_map = {
            str(code or "").strip(): cls._safe_float(total)
            for code, total in (
                db.session.query(Entrada.codigo_item, func.coalesce(func.sum(Entrada.quantidade), 0.0))
                .filter(Entrada.codigo_item.in_(codes))
                .group_by(Entrada.codigo_item)
                .all()
            )
            if str(code or "").strip()
        }
        saidas_map = {
            str(code or "").strip(): cls._safe_float(total)
            for code, total in (
                db.session.query(Saida.codigo_item, func.coalesce(func.sum(Saida.quantidade), 0.0))
                .filter(Saida.codigo_item.in_(codes))
                .group_by(Saida.codigo_item)
                .all()
            )
            if str(code or "").strip()
        }
        eventos_map = {
            str(code or "").strip(): cls._safe_float(total)
            for code, total in (
                db.session.query(InventarioEvento.codigo_item, func.coalesce(func.sum(InventarioEvento.quantidade), 0.0))
                .filter(InventarioEvento.codigo_item.in_(codes))
                .group_by(InventarioEvento.codigo_item)
                .all()
            )
            if str(code or "").strip()
        }
        ledger_map = {
            str(code or "").strip(): cls._safe_float(total)
            for code, total in (
                db.session.query(StockMovement.product_id, func.coalesce(func.sum(StockMovement.quantity_base), 0.0))
                .filter(StockMovement.product_id.in_(codes))
                .group_by(StockMovement.product_id)
                .all()
            )
            if str(code or "").strip()
        }
        cache_map = {
            str(row.product_id or "").strip(): cls._safe_float(row.quantity_base)
            for row in (
                db.session.query(StockBalance)
                .filter(StockBalance.product_id.in_(codes))
                .all()
            )
            if str(row.product_id or "").strip()
        }

        sections: dict[str, list[dict[str, Any]]] = {
            "confirmed": [],
            "probable": [],
            "future": [],
        }

        for item in scoped_items:
            code = str(item.codigo_item or "").strip()
            if not code:
                continue

            summary = summary_by_code.get(code, {})
            saldo_sistema = cls._safe_float(summary.get("saldo_base") or item.get_saldo_fisico_total())
            valor_total = cls._safe_float(summary.get("valor_total"))
            categoria = str(item.categoria or "Sem categoria").strip() or "Sem categoria"
            packaging_type = str(getattr(item, "tipo_embalagem_novo", None) or getattr(item, "tipo_embalagem", None) or "").strip() or "-"
            packaging_factor = cls._safe_float(resolve_packaging_factor(item))
            legacy_balance = cls._safe_float(entradas_map.get(code)) - cls._safe_float(saidas_map.get(code)) + cls._safe_float(eventos_map.get(code))
            ledger_balance = cls._safe_float(ledger_map.get(code))
            cache_balance = cls._safe_float(cache_map.get(code, ledger_balance))

            if ignore_packaging_metadata_for_stock(item):
                if (
                    abs(legacy_balance - ledger_balance) > cls.ANALYSIS_TOLERANCE
                    or abs(ledger_balance - cache_balance) > cls.ANALYSIS_TOLERANCE
                ):
                    sections["confirmed"].append(
                        {
                            "codigo": code,
                            "descricao": item.descricao or code,
                            "categoria": categoria,
                            "saldo_sistema": saldo_sistema,
                            "valor_total": valor_total,
                            "signal": "Toolkit/jogo com saldo inflado por multiplicacao interna",
                            "reference": f"Legado {legacy_balance:g} | Ledger {ledger_balance:g} | Cache {cache_balance:g}",
                            "impact": "Infla saldo fisico e valor monetario do relatório.",
                        }
                    )
                else:
                    sections["future"].append(
                        {
                            "codigo": code,
                            "descricao": item.descricao or code,
                            "categoria": categoria,
                            "saldo_sistema": saldo_sistema,
                            "valor_total": valor_total,
                            "signal": "Toolkit alinhado, mas ainda com metadado de embalagem residual",
                            "reference": f"Tipo {packaging_type} | fator {packaging_factor:g} | embalagem ignorada pelo estoque",
                            "impact": "Pode voltar a inflar em rebuild legado ou leitura financeira futura se o fluxo errado for reativado.",
                        }
                    )

            if packaging_type != "-" and packaging_factor <= 0 and not ignore_packaging_metadata_for_stock(item):
                sections["probable"].append(
                    {
                        "codigo": code,
                        "descricao": item.descricao or code,
                        "categoria": categoria,
                        "saldo_sistema": saldo_sistema,
                        "valor_total": valor_total,
                        "signal": "Embalagem sem fator confiavel de conversao",
                        "reference": f"Tipo {packaging_type} | fator resolvido {packaging_factor:g}",
                        "impact": "Pode distorcer saldo fisico, decomposicao visual e totalizacao por valor.",
                    }
                )

            if bool(summary.get("has_price_alert")):
                sections["probable"].append(
                    {
                        "codigo": code,
                        "descricao": item.descricao or code,
                        "categoria": categoria,
                        "saldo_sistema": saldo_sistema,
                        "valor_total": valor_total,
                        "signal": "Preco base com sinal de embalagem nao normalizada",
                        "reference": str(summary.get("price_alert") or summary.get("origem_preco") or "Sem detalhe"),
                        "impact": "Pode inflar ou subestimar o valor em estoque no ranking monetario.",
                    }
                )

        for key in sections:
            sections[key] = sorted(
                sections[key],
                key=lambda row: (
                    -cls._safe_float(row.get("valor_total")),
                    -cls._safe_float(row.get("saldo_sistema")),
                    str(row.get("descricao") or "").lower(),
                ),
            )

        return sections

    @classmethod
    def _merge_row(cls, ws, row: int, max_col: int, value: str, *, font: Font, fill: PatternFill | None = None, alignment: Alignment | None = None):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max_col)
        cell = ws.cell(row=row, column=1)
        cell.value = value
        cell.font = font
        if fill is not None:
            cell.fill = fill
        cell.alignment = alignment or Alignment(horizontal="left", vertical="center")

    @classmethod
    def _style_section_banner(cls, ws, *, row: int, max_col: int, title: str, color: str) -> None:
        cls._merge_row(
            ws,
            row,
            max_col,
            title,
            font=Font(size=11, bold=True, color=cls.COLOR_WHITE),
            fill=PatternFill(start_color=color, end_color=color, fill_type="solid"),
            alignment=Alignment(horizontal="center", vertical="center"),
        )
        ws.row_dimensions[row].height = 20

    @classmethod
    def _apply_report_header(cls, ws, *, title: str, scope_label: str, max_col: int, meta_lines: list[str] | None = None) -> int:
        ws.sheet_view.showGridLines = False
        row = 1
        cls._merge_row(
            ws,
            row,
            max_col,
            title,
            font=Font(size=14, bold=True, color=cls.COLOR_WHITE),
            fill=PatternFill(start_color=cls.COLOR_HEADER, end_color=cls.COLOR_HEADER, fill_type="solid"),
            alignment=Alignment(horizontal="center", vertical="center"),
        )
        ws.row_dimensions[row].height = 24
        row += 1

        for line in get_company_header_lines():
            cls._merge_row(
                ws,
                row,
                max_col,
                line,
                font=Font(size=10, bold=True, color=cls.COLOR_HEADER_DARK),
                alignment=Alignment(horizontal="center", vertical="center"),
            )
            row += 1

        for line in [
            f"Escopo: {scope_label}",
            f"Gerado em: {datetime.now().strftime('%d/%m/%Y às %H:%M')}",
            *((meta_lines or [])),
        ]:
            cls._merge_row(
                ws,
                row,
                max_col,
                line,
                font=Font(size=10, color=cls.COLOR_MUTED),
                alignment=Alignment(horizontal="center", vertical="center"),
            )
            row += 1

        return row + 1

    @classmethod
    def _style_table_header(cls, ws, row: int, headers: list[str]) -> None:
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col_num)
            cell.value = header
            cell.font = Font(bold=True, color=cls.COLOR_WHITE, size=10)
            cell.fill = PatternFill(start_color=cls.COLOR_HEADER, end_color=cls.COLOR_HEADER, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = cls.BORDER_THIN

    @classmethod
    def _style_table_row(cls, ws, row: int, max_col: int, *, alert: bool = False) -> None:
        fill_color = cls.COLOR_ALERT if alert else (cls.COLOR_LIGHT_ALT if row % 2 == 0 else cls.COLOR_WHITE)
        fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = cls.BORDER_THIN
            cell.fill = fill
            if col == 1:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    @classmethod
    def _set_widths(cls, ws, widths: dict[str, float]) -> None:
        for column, width in widths.items():
            ws.column_dimensions[column].width = width

    @classmethod
    def _add_ranking_chart(
        cls,
        ws,
        *,
        title: str,
        header_row: int,
        category_col: int,
        value_col: int,
        data_start_row: int,
        data_end_row: int,
        anchor: str,
        color: str,
        value_axis_title: str,
    ) -> None:
        if data_end_row < data_start_row:
            return

        data = Reference(ws, min_col=value_col, min_row=header_row, max_row=data_end_row)
        categories = Reference(ws, min_col=category_col, min_row=data_start_row, max_row=data_end_row)

        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = title
        chart.grouping = "clustered"
        chart.overlap = 0
        chart.height = 6.8
        chart.width = 10.8
        chart.y_axis.title = "Itens"
        chart.x_axis.title = value_axis_title
        chart.legend = None
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)

        if chart.series:
            series = chart.series[0]
            series.graphicalProperties.solidFill = color
            series.graphicalProperties.line.solidFill = color

        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
        ws.add_chart(chart, anchor)

    @classmethod
    def _add_timeline_chart(cls, ws, *, header_row: int, data_start_row: int, data_end_row: int, anchor: str) -> None:
        if data_end_row < data_start_row:
            return

        data = Reference(ws, min_col=2, max_col=4, min_row=header_row, max_row=data_end_row)
        categories = Reference(ws, min_col=1, min_row=data_start_row, max_row=data_end_row)

        chart = LineChart()
        chart.style = 2
        chart.title = "Timeline de movimentações"
        chart.y_axis.title = "Eventos"
        chart.x_axis.title = "Período"
        chart.height = 7.2
        chart.width = 13.5
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)

        colors = [cls.COLOR_SUCCESS, cls.COLOR_DANGER, cls.COLOR_HEADER_DARK]
        for index, series in enumerate(chart.series):
            color = colors[index] if index < len(colors) else cls.COLOR_HEADER
            series.graphicalProperties.line.solidFill = color
            series.graphicalProperties.solidFill = color

        ws.add_chart(chart, anchor)

    @classmethod
    def _create_executive_summary_sheet(
        cls,
        wb: Workbook,
        category_analysis: dict[str, dict[str, Any]],
        items: list[dict[str, Any]],
        timeline: list[dict[str, Any]],
        scope_label: str,
    ) -> None:
        ws = wb.create_sheet("Resumo Executivo")
        total_valor_real = round(sum(cls._safe_float(item["valor_total_real"]) for item in items), 2)
        total_valor_estimado = round(sum(cls._safe_float(item["valor_total_estimado"]) for item in items), 2)
        total_alertas = sum(1 for item in items if item["has_price_alert"])
        total_itens_com_saldo = sum(1 for item in items if cls._safe_float(item["saldo_base"]) > 0)
        ultima_movimentacao = max((item.get("ultima_movimentacao_dt") for item in items if item.get("ultima_movimentacao_dt")), default=None)

        row = cls._apply_report_header(
            ws,
            title="RELATORIO DE ESCOPO - RESUMO EXECUTIVO",
            scope_label=scope_label,
            max_col=9,
            meta_lines=[
                "Leitura consolidada com cabecalho operacional, timeline e alertas de precificacao.",
                "Obs.: o consolidado evita somar quantidades fisicas heterogeneas entre itens de unidades diferentes.",
            ],
        )

        metrics = [
            ("Itens no escopo", str(len(items))),
            ("Categorias ativas", str(len(category_analysis))),
            ("Itens com saldo", str(total_itens_com_saldo)),
            ("Valor documentado", f"R$ {total_valor_real:,.2f}"),
            ("Valor estimado", f"R$ {total_valor_estimado:,.2f}"),
            ("Alertas de preco", str(total_alertas)),
            ("Ultima movimentacao", cls._format_datetime(ultima_movimentacao)),
        ]

        panel_title_row = row
        cls._merge_row(
            ws,
            panel_title_row,
            9,
            "PAINEL EXECUTIVO",
            font=Font(size=11, bold=True, color=cls.COLOR_HEADER_DARK),
            fill=PatternFill(start_color=cls.COLOR_PANEL, end_color=cls.COLOR_PANEL, fill_type="solid"),
        )
        row += 1
        for label, value in metrics:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
            ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=6)
            ws.merge_cells(start_row=row, start_column=7, end_row=row, end_column=9)
            left = ws.cell(row=row, column=1)
            center = ws.cell(row=row, column=4)
            right = ws.cell(row=row, column=7)
            left.value = label
            center.value = value
            right.value = "" if label != "Alertas de preco" else "Linhas em alerta aparecem destacadas no detalhamento."
            left.font = Font(bold=True, color=cls.COLOR_TEXT)
            center.font = Font(bold=True, color=cls.COLOR_HEADER_DARK)
            right.font = Font(size=9, color=cls.COLOR_MUTED)
            for cell in (left, center, right):
                cell.border = cls.BORDER_THIN
                cell.fill = PatternFill(start_color=cls.COLOR_LIGHT if row % 2 else cls.COLOR_LIGHT_ALT, end_color=cls.COLOR_LIGHT if row % 2 else cls.COLOR_LIGHT_ALT, fill_type="solid")
                cell.alignment = Alignment(horizontal="left", vertical="center")
            row += 1

        row += 1
        headers = [
            "Categoria",
            "Itens",
            "Itens com saldo",
            "Mov. entradas",
            "Mov. saidas",
            "Valor documentado",
            "Valor estimado",
            "Alertas",
            "Ultima movimentacao",
        ]
        header_row = row
        cls._style_table_header(ws, header_row, headers)
        row += 1

        for category_name, data in sorted(category_analysis.items(), key=lambda row_data: (-(row_data[1]["valor_real"] + row_data[1]["valor_estimado"]), row_data[0])):
            ws.cell(row=row, column=1, value=category_name)
            ws.cell(row=row, column=2, value=int(data["total_itens"] or 0))
            ws.cell(row=row, column=3, value=int(data["itens_com_saldo"] or 0))
            ws.cell(row=row, column=4, value=int(data["entrada_eventos"] or 0))
            ws.cell(row=row, column=5, value=int(data["saida_eventos"] or 0))
            ws.cell(row=row, column=6, value=round(cls._safe_float(data["valor_real"]), 2)).number_format = '"R$" #,##0.00'
            ws.cell(row=row, column=7, value=round(cls._safe_float(data["valor_estimado"]), 2)).number_format = '"R$" #,##0.00'
            ws.cell(row=row, column=8, value=int(data["itens_com_alerta"] or 0))
            ws.cell(row=row, column=9, value=cls._format_datetime(data.get("ultima_movimentacao_dt")))
            cls._style_table_row(ws, row, 9, alert=bool(data.get("itens_com_alerta")))
            for col in range(2, 9):
                ws.cell(row=row, column=col).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=row, column=9).alignment = Alignment(horizontal="left", vertical="center")
            row += 1

        note_row = row + 1
        recent_window = timeline[-3:] if len(timeline) >= 3 else timeline
        timeline_note = " | ".join(
            f"{entry['period_label']}: {entry['movimentacoes']} movimentos"
            for entry in recent_window
        ) or "Sem movimentacoes no periodo auditado."
        cls._merge_row(
            ws,
            note_row,
            9,
            f"TIMELINE RECENTE: {timeline_note}",
            font=Font(size=10, italic=True, color=cls.COLOR_MUTED),
        )

        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = f"A{header_row}:I{row - 1}"
        cls._set_widths(ws, {
            "A": 28,
            "B": 10,
            "C": 14,
            "D": 13,
            "E": 13,
            "F": 18,
            "G": 18,
            "H": 10,
            "I": 22,
        })

    @classmethod
    def _create_timeline_sheet(cls, wb: Workbook, timeline: list[dict[str, Any]], scope_label: str) -> None:
        ws = wb.create_sheet("Timeline")
        row = cls._apply_report_header(
            ws,
            title="RELATORIO DE ESCOPO - TIMELINE OPERACIONAL",
            scope_label=scope_label,
            max_col=7,
            meta_lines=[
                "Timeline mensal de eventos de entrada e saida dentro do escopo selecionado.",
                "Os valores abaixo representam contagem de movimentos, itens e categorias, evitando mistura de unidades fisicas.",
            ],
        )

        headers = [
            "Periodo",
            "Entradas",
            "Saidas",
            "Mov. totais",
            "Itens movimentados",
            "Categorias ativas",
            "Ultimo evento do periodo",
        ]
        header_row = row
        cls._style_table_header(ws, header_row, headers)
        row += 1

        for entry in timeline:
            ws.cell(row=row, column=1, value=entry["period_label"])
            ws.cell(row=row, column=2, value=int(entry["entradas"] or 0))
            ws.cell(row=row, column=3, value=int(entry["saidas"] or 0))
            ws.cell(row=row, column=4, value=int(entry["movimentacoes"] or 0))
            ws.cell(row=row, column=5, value=int(entry["itens_movimentados"] or 0))
            ws.cell(row=row, column=6, value=int(entry["categorias_ativas"] or 0))
            ws.cell(row=row, column=7, value=entry["ultimo_evento_label"])
            cls._style_table_row(ws, row, 7)
            for col in range(2, 7):
                ws.cell(row=row, column=col).alignment = Alignment(horizontal="center", vertical="center")
            row += 1

        cls._add_timeline_chart(
            ws,
            header_row=header_row,
            data_start_row=header_row + 1,
            data_end_row=row - 1,
            anchor="I3",
        )

        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = f"A{header_row}:G{max(row - 1, header_row)}"
        cls._set_widths(ws, {
            "A": 12,
            "B": 10,
            "C": 10,
            "D": 12,
            "E": 18,
            "F": 16,
            "G": 42,
            "I": 2,
            "J": 14,
            "K": 14,
            "L": 14,
        })

    @classmethod
    def _create_brand_analysis_sheet(cls, wb: Workbook, brand_analysis: dict[str, dict[str, Any]], scope_label: str) -> None:
        ws = wb.create_sheet("Marcas e Cobertura")
        total_value = round(sum(cls._safe_float(row["valor_total"]) for row in brand_analysis.values()), 2)
        row = cls._apply_report_header(
            ws,
            title="RELATORIO DE ESCOPO - MARCAS E COBERTURA",
            scope_label=scope_label,
            max_col=6,
            meta_lines=[
                f"Valor total analisado nesta aba: R$ {total_value:,.2f}",
                "A cobertura por marca ajuda a identificar concentracao de valor e lacunas de cadastro.",
            ],
        )

        headers = ["Categoria", "Marca", "Itens", "Itens com saldo", "Valor total", "Alertas"]
        header_row = row
        cls._style_table_header(ws, header_row, headers)
        row += 1

        for data in sorted(brand_analysis.values(), key=lambda current: (-current["valor_total"], current["categoria"], current["marca"])):
            ws.cell(row=row, column=1, value=data["categoria"])
            ws.cell(row=row, column=2, value=data["marca"])
            ws.cell(row=row, column=3, value=int(data["total_itens"] or 0))
            ws.cell(row=row, column=4, value=int(data["itens_com_saldo"] or 0))
            ws.cell(row=row, column=5, value=round(cls._safe_float(data["valor_total"]), 2)).number_format = '"R$" #,##0.00'
            ws.cell(row=row, column=6, value=int(data["itens_com_alerta"] or 0))
            cls._style_table_row(ws, row, 6, alert=bool(data.get("itens_com_alerta")))
            for col in range(3, 7):
                ws.cell(row=row, column=col).alignment = Alignment(horizontal="center", vertical="center")
            row += 1

        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = f"A{header_row}:F{row - 1}"
        cls._set_widths(ws, {
            "A": 28,
            "B": 24,
            "C": 10,
            "D": 15,
            "E": 16,
            "F": 10,
        })

    @classmethod
    def _create_rankings_sheet(cls, wb: Workbook, rankings: dict[str, list[dict[str, Any]]], scope_label: str) -> None:
        ws = wb.create_sheet("Rankings")
        row = cls._apply_report_header(
            ws,
            title="RELATORIO DE ESCOPO - RANKINGS DE LEITURA",
            scope_label=scope_label,
            max_col=14,
            meta_lines=[
                "Rankings priorizam comparacoes monetarias e frequencia de eventos, evitando mistura de quantidades fisicas incompatíveis.",
            ],
        )

        sections = [
            {
                "title": "TOP 10 MAIOR VALOR EM ESTOQUE",
                "color": cls.COLOR_DANGER,
                "rows": rankings.get("mais_caros") or [],
                "headers": ["#", "Codigo", "Descricao", "Categoria", "Marca", "Valor total"],
                "chart_title": "Itens com maior valor em estoque",
                "chart_value_title": "Valor em estoque (R$)",
                "type": "valor",
            },
            {
                "title": "TOP 10 MAIOR FREQUENCIA DE SAIDAS",
                "color": cls.COLOR_SUCCESS,
                "rows": rankings.get("mais_usados") or [],
                "headers": ["#", "Codigo", "Descricao", "Categoria", "Marca", "Saidas"],
                "chart_title": "Itens com mais saidas",
                "chart_value_title": "Saidas (eventos)",
                "type": "saida_eventos",
            },
            {
                "title": "TOP 10 MAIOR MOVIMENTACAO TOTAL",
                "color": cls.COLOR_HEADER_DARK,
                "rows": rankings.get("mais_movimentados") or [],
                "headers": ["#", "Codigo", "Descricao", "Categoria", "Marca", "Mov. totais"],
                "chart_title": "Itens mais movimentados",
                "chart_value_title": "Movimentos (eventos)",
                "type": "total_movimentacoes",
            },
            {
                "title": "TOP 10 ALERTAS DE PRECIFICACAO",
                "color": cls.COLOR_WARNING,
                "rows": rankings.get("com_alerta") or [],
                "headers": ["#", "Codigo", "Descricao", "Categoria", "Valor total", "Sinalizacao"],
                "chart_title": None,
                "chart_value_title": None,
                "type": "alerta",
            },
        ]

        current_row = row
        for section in sections:
            section_start_row = current_row
            cls._style_section_banner(
                ws,
                row=current_row,
                max_col=14,
                title=section["title"],
                color=section["color"],
            )
            current_row += 1

            header_row = current_row
            cls._style_table_header(ws, header_row, section["headers"])
            current_row += 1

            data_start_row = current_row
            for index, item in enumerate(section["rows"], 1):
                ws.cell(row=current_row, column=1, value=index)
                ws.cell(row=current_row, column=2, value=item["codigo"])
                ws.cell(row=current_row, column=3, value=item["descricao"])
                ws.cell(row=current_row, column=4, value=item["categoria"])
                if section["type"] == "valor":
                    ws.cell(row=current_row, column=5, value=item["marca"] or "N/D")
                    ws.cell(row=current_row, column=6, value=round(cls._safe_float(item["valor_total"]), 2)).number_format = '"R$" #,##0.00'
                elif section["type"] == "saida_eventos":
                    ws.cell(row=current_row, column=5, value=item["marca"] or "N/D")
                    ws.cell(row=current_row, column=6, value=int(item["saida_eventos"] or 0))
                elif section["type"] == "total_movimentacoes":
                    ws.cell(row=current_row, column=5, value=item["marca"] or "N/D")
                    ws.cell(row=current_row, column=6, value=int(item["total_movimentacoes"] or 0))
                else:
                    ws.cell(row=current_row, column=5, value=round(cls._safe_float(item["valor_total"]), 2)).number_format = '"R$" #,##0.00'
                    ws.cell(row=current_row, column=6, value=item["price_alert"] or "-")
                cls._style_table_row(ws, current_row, 6, alert=item["has_price_alert"])
                current_row += 1

            data_end_row = current_row - 1
            if section["chart_title"]:
                cls._add_ranking_chart(
                    ws,
                    title=section["chart_title"],
                    header_row=header_row,
                    category_col=3,
                    value_col=6,
                    data_start_row=data_start_row,
                    data_end_row=data_end_row,
                    anchor=f"H{section_start_row}",
                    color=section["color"],
                    value_axis_title=section["chart_value_title"],
                )
                current_row = max(current_row + 2, section_start_row + 18)
            else:
                current_row += 2

        ws.freeze_panes = f"A{row + 1}"
        cls._set_widths(ws, {
            "A": 5,
            "B": 16,
            "C": 44,
            "D": 24,
            "E": 18,
            "F": 22,
            "G": 3,
            "H": 3,
            "I": 14,
            "J": 14,
            "K": 14,
            "L": 14,
            "M": 14,
            "N": 14,
        })

    @classmethod
    def _create_inconsistency_sheet(cls, wb: Workbook, analysis: dict[str, list[dict[str, Any]]], scope_label: str) -> None:
        ws = wb.create_sheet("Inconsistencias")
        confirmed = analysis.get("confirmed") or []
        probable = analysis.get("probable") or []
        future = analysis.get("future") or []

        row = cls._apply_report_header(
            ws,
            title="RELATORIO DE ESCOPO - INCONSISTENCIAS E RISCOS",
            scope_label=scope_label,
            max_col=9,
            meta_lines=[
                "Esta aba separa o que ja esta inconsistente, o que parece inconsistente e o que pode voltar a quebrar futuramente.",
                "Ela usa sinais do proprio GALINT: toolkit divergente, embalagem sem fator confiavel e alertas de precificacao.",
            ],
        )

        summary_metrics = [
            ("Inconsistencias confirmadas", len(confirmed), cls.COLOR_DANGER),
            ("Provaveis inconsistencias", len(probable), cls.COLOR_WARNING),
            ("Riscos futuros", len(future), cls.COLOR_HEADER_DARK),
        ]
        for label, value, color in summary_metrics:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=9)
            left = ws.cell(row=row, column=1)
            right = ws.cell(row=row, column=5)
            left.value = label
            right.value = value
            left.font = Font(bold=True, color=cls.COLOR_TEXT)
            right.font = Font(bold=True, color=color)
            for cell in (left, right):
                cell.border = cls.BORDER_THIN
                cell.fill = PatternFill(start_color=cls.COLOR_LIGHT_ALT, end_color=cls.COLOR_LIGHT_ALT, fill_type="solid")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            row += 1

        row += 1
        section_specs = [
            ("INCONSISTENCIAS CONFIRMADAS", confirmed, cls.COLOR_DANGER),
            ("PROVAVEIS INCONSISTENCIAS", probable, cls.COLOR_WARNING),
            ("RISCOS FUTUROS", future, cls.COLOR_HEADER_DARK),
        ]
        headers = ["#", "Codigo", "Descricao", "Categoria", "Saldo sistema", "Valor estoque", "Sinal", "Leitura de apoio", "Impacto"]

        for title, rows, color in section_specs:
            cls._style_section_banner(ws, row=row, max_col=9, title=title, color=color)
            row += 1
            header_row = row
            cls._style_table_header(ws, header_row, headers)
            row += 1

            if not rows:
                cls._merge_row(
                    ws,
                    row,
                    9,
                    "Nenhum item identificado nesta faixa de analise.",
                    font=Font(size=10, italic=True, color=cls.COLOR_MUTED),
                    fill=PatternFill(start_color=cls.COLOR_LIGHT, end_color=cls.COLOR_LIGHT, fill_type="solid"),
                    alignment=Alignment(horizontal="center", vertical="center"),
                )
                for col in range(1, 10):
                    ws.cell(row=row, column=col).border = cls.BORDER_THIN
                row += 2
                continue

            for index, entry in enumerate(rows, 1):
                ws.cell(row=row, column=1, value=index)
                ws.cell(row=row, column=2, value=entry["codigo"])
                ws.cell(row=row, column=3, value=entry["descricao"])
                ws.cell(row=row, column=4, value=entry["categoria"])
                ws.cell(row=row, column=5, value=round(cls._safe_float(entry["saldo_sistema"]), 2))
                ws.cell(row=row, column=6, value=round(cls._safe_float(entry["valor_total"]), 2)).number_format = '"R$" #,##0.00'
                ws.cell(row=row, column=7, value=entry["signal"])
                ws.cell(row=row, column=8, value=entry["reference"])
                ws.cell(row=row, column=9, value=entry["impact"])
                cls._style_table_row(ws, row, 9, alert=title != "RISCOS FUTUROS")
                ws.cell(row=row, column=1).alignment = Alignment(horizontal="center", vertical="center")
                ws.cell(row=row, column=5).alignment = Alignment(horizontal="center", vertical="center")
                ws.cell(row=row, column=6).alignment = Alignment(horizontal="center", vertical="center")
                row += 1

            row += 2

        ws.freeze_panes = "A8"
        cls._set_widths(ws, {
            "A": 5,
            "B": 16,
            "C": 38,
            "D": 22,
            "E": 14,
            "F": 16,
            "G": 32,
            "H": 36,
            "I": 42,
        })

    @classmethod
    def _create_detailed_sheet(cls, wb: Workbook, items: list[dict[str, Any]], scope_label: str) -> None:
        ws = wb.create_sheet("Detalhamento")
        row = cls._apply_report_header(
            ws,
            title="RELATORIO DE ESCOPO - DETALHAMENTO",
            scope_label=scope_label,
            max_col=15,
            meta_lines=[
                "Detalhamento por item com saldo fisico exibido, eventos, origem de preco e sinalizacao de auditoria.",
            ],
        )

        headers = [
            "Codigo",
            "Descricao",
            "Categoria",
            "Marca",
            "Saldo fisico",
            "Saldo base",
            "Entradas",
            "Saidas",
            "Ult. entrada",
            "Ult. saida",
            "Preco base",
            "Unidade preco",
            "Valor estoque",
            "Origem",
            "Sinalizacao",
        ]
        header_row = row
        cls._style_table_header(ws, header_row, headers)
        row += 1

        for item in sorted(items, key=lambda current: (current["categoria"], current["descricao"])):
            ws.cell(row=row, column=1, value=item["codigo"])
            ws.cell(row=row, column=2, value=item["descricao"])
            ws.cell(row=row, column=3, value=item["categoria"])
            ws.cell(row=row, column=4, value=item["marca"] or "N/D")
            ws.cell(row=row, column=5, value=item["saldo_display"])
            ws.cell(row=row, column=6, value=round(cls._safe_float(item["saldo_base"]), 4)).number_format = "#,##0.0000"
            ws.cell(row=row, column=7, value=int(item["entrada_eventos"] or 0))
            ws.cell(row=row, column=8, value=int(item["saida_eventos"] or 0))
            ws.cell(row=row, column=9, value=item["ultima_entrada_label"])
            ws.cell(row=row, column=10, value=item["ultima_saida_label"])
            ws.cell(row=row, column=11, value=round(cls._safe_float(item["preco_unitario_base"]), 6) if item["preco_unitario_base"] else 0).number_format = '"R$" #,##0.000000'
            ws.cell(row=row, column=12, value=(item["preco_unidade"] or "-").upper())
            ws.cell(row=row, column=13, value=round(cls._safe_float(item["valor_total"]), 2)).number_format = '"R$" #,##0.00'
            ws.cell(row=row, column=14, value=item["origem_preco"])
            ws.cell(row=row, column=15, value=item["price_alert"] or "-")
            cls._style_table_row(ws, row, 15, alert=item["has_price_alert"])
            for col in (6, 7, 8, 11, 13):
                ws.cell(row=row, column=col).alignment = Alignment(horizontal="center", vertical="center")
            row += 1

        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = f"A{header_row}:O{row - 1}"
        cls._set_widths(ws, {
            "A": 16,
            "B": 40,
            "C": 24,
            "D": 18,
            "E": 28,
            "F": 12,
            "G": 10,
            "H": 10,
            "I": 18,
            "J": 18,
            "K": 16,
            "L": 14,
            "M": 16,
            "N": 30,
            "O": 34,
        })
