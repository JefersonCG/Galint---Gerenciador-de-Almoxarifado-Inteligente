from __future__ import annotations

"""Serviço responsável pelo registro do Inventário de Materiais Danificados."""

from collections.abc import Iterable
from datetime import datetime
import re
from pathlib import Path
from typing import Any

from docx import Document
from flask import current_app
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..extensions import db
from ..models import DescarteAutorizacao, MaterialInventario, Saida
from ..utils.report_branding import get_company_header_lines, get_company_header_text


class MaterialInventoryService:
    """Gerencia registros oficiais de materiais devolvidos danificados."""

    def list_reports(self, limit: int = 200) -> list[dict[str, Any]]:
        registros = (
            MaterialInventario.query.order_by(MaterialInventario.data_registro.desc()).limit(limit).all()
        )
        return [self._to_dict(entry) for entry in registros]

    def list_available_saidas(self, limit: int = 20) -> list[dict[str, Any]]:
        logged = {
            row[0]
            for row in db.session.query(MaterialInventario.saida_id)
            .filter(MaterialInventario.saida_id.isnot(None))
            .all()
            if row[0]
        }
        query = Saida.query.order_by(Saida.data_saida.desc())
        if logged:
            query = query.filter(~Saida.id_saida.in_(logged))
        registros = query.limit(limit).all()
        return [self._map_saida(saida) for saida in registros]

    def record_damage(self, saida_id: int, quantidade: int, descricao: str | None, observacoes: str | None) -> MaterialInventario:
        saida = Saida.query.get(saida_id)
        if not saida:
            raise ValueError("Saída não encontrada")
        quantidade = int(quantidade)
        if quantidade <= 0:
            raise ValueError("Informe quantidade positiva")
        if quantidade > saida.quantidade:
            raise ValueError("Quantidade excede o registro da saída")
        if MaterialInventario.query.filter_by(saida_id=saida_id).first():
            raise ValueError("Esta saída já foi registrada como danificada")
        registro = MaterialInventario(
            saida_id=saida.id_saida,
            codigo_item=saida.codigo_item,
            matricula=saida.matricula,
            responsavel_nome=saida.usuario.nome if saida.usuario else None,
            quantidade=quantidade,
            descricao=(descricao or "").strip() or "Material devolvido danificado",
            observacoes=(observacoes or "").strip() or None,
        )
        db.session.add(registro)
        db.session.commit()
        return registro

    def list_authorizations(self, limit: int = 10) -> list[dict[str, Any]]:
        autorizacoes = (
            DescarteAutorizacao.query.order_by(DescarteAutorizacao.data_autorizacao.desc()).limit(limit).all()
        )
        return [
            {
                "id": auth.id,
                "gerente": auth.gerente_nome,
                "comentario": auth.comentario,
                "arquivo": auth.arquivo_relatorio,
                "data": auth.data_autorizacao,
            }
            for auth in autorizacoes
        ]

    def authorize_disposal(
        self,
        gerente_nome: str,
        comentario: str | None,
        relatorio_html: str | None,
        entries: Iterable[dict[str, Any]],
    ) -> DescarteAutorizacao:
        gerente = (gerente_nome or "").strip()
        if not gerente:
            raise ValueError("Informe o gerente responsável pela autorização")

        now = datetime.utcnow()
        workbook = self._build_report_workbook(gerente, comentario, relatorio_html, entries, now)
        reports_dir = self._ensure_reports_dir()
        file_name = f"produtos-avariados-{now.strftime('%Y%m%d%H%M%S')}.xlsx"
        file_path = reports_dir / file_name
        workbook.save(file_path)

        registro = DescarteAutorizacao(
            gerente_nome=gerente,
            comentario=(comentario or "").strip() or None,
            conteudo=relatorio_html,
            arquivo_relatorio=file_name,
        )
        db.session.add(registro)
        db.session.commit()
        return registro

    def _ensure_reports_dir(self) -> Path:
        reports_path = Path(current_app.instance_path) / "reports"
        reports_path.mkdir(parents=True, exist_ok=True)
        return reports_path

    def _strip_html(self, raw: str | None) -> str:
        if not raw:
            return ""
        text = re.sub(r"<[^>]+>", "", raw)
        return text.strip()

    def _format_codigo_barra(self, value: object) -> str:
        if value is None:
            return "N/D"
        text = str(value).strip()
        if not text:
            return "N/D"
        return text[-4:] if len(text) >= 4 else text

    def _build_report_workbook(
        self,
        gerente: str,
        comentario: str | None,
        relatorio_html: str | None,
        entries: Iterable[dict[str, Any]],
        timestamp: datetime,
    ) -> Workbook:
        wb = Workbook()
        ws = wb.active
        ws.title = "Produtos Avariados"

        total_columns = 6
        end_column = get_column_letter(total_columns)

        # Cabeçalho da empresa (padronizado)
        current_row = 1
        company_lines = get_company_header_lines()
        if company_lines:
            for idx, line in enumerate(company_lines):
                ws.merge_cells(f"A{current_row}:{end_column}{current_row}")
                cell = ws[f"A{current_row}"]
                cell.value = line
                cell.font = Font(bold=(idx == 0), size=18 if idx == 0 else 11)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                ws.row_dimensions[current_row].height = 28 if idx == 0 else 18
                current_row += 1

        # Título do relatório
        ws.merge_cells(f"A{current_row}:{end_column}{current_row}")
        ws[f"A{current_row}"] = "Relatório de Produtos Avariados"
        ws[f"A{current_row}"].font = Font(bold=True, size=16)
        ws[f"A{current_row}"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[current_row].height = 26
        current_row += 2

        # Informações do relatório
        ws[f"A{current_row}"] = f"Data de geração: {timestamp.strftime('%d/%m/%Y %H:%M')} UTC"
        current_row += 1
        ws[f"A{current_row}"] = f"Gerente responsável: {gerente}"
        current_row += 1
        
        if comentario:
            ws[f"A{current_row}"] = f"Observações adicionais: {comentario}"
            current_row += 1

        cleaned = self._strip_html(relatorio_html)
        if cleaned:
            ws[f"A{current_row}"] = "Resumo do relatório:"
            current_row += 1
            ws[f"A{current_row}"] = cleaned
            ws[f"A{current_row}"].alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[current_row].height = max(15 * len(cleaned.split('\n')), 30)
            current_row += 1

        # Espaço antes da tabela
        current_row += 1

        # Cabeçalhos da tabela
        headers = ["Código", "Item", "Qnt", "Responsável", "Saída", "Observações"]
        header_row = current_row
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        header_fill = PatternFill("solid", fgColor="FDE047")

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col_idx)
            cell.value = header
            cell.fill = header_fill
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        current_row += 1

        # Dados da tabela
        for entry in entries:
            row_cells = [
                self._format_codigo_barra(entry.get("codigo")),
                entry.get("descricao", "-") or "-",
                str(entry.get("quantidade", 0)),
                entry.get("responsavel", "-") or "-",
                entry.get("saida_data").strftime('%d/%m/%Y %H:%M') if entry.get("saida_data") else "-",
                entry.get("observacoes", "-") or "-",
            ]
            
            for col_idx, value in enumerate(row_cells, start=1):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.value = value
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            
            current_row += 1

        # Assinatura
        current_row += 2
        ws[f"A{current_row}"] = "Assinatura digital do gerente: ________________________________"

        # Ajustar larguras das colunas
        column_widths = [8, 45, 8, 25, 20, 40]
        for idx, width in enumerate(column_widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width

        return wb

    def _build_report_document(
        self,
        gerente: str,
        comentario: str | None,
        relatorio_html: str | None,
        entries: Iterable[dict[str, Any]],
        timestamp: datetime,
    ) -> Document:
        doc = Document()
        section = doc.sections[0]
        header = section.header
        header_paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        header_paragraph.text = get_company_header_text()
        header_paragraph.alignment = 1

        doc.add_heading("Produtos Avariados / Refugo", level=1)
        doc.add_paragraph(f"Data de geração: {timestamp.strftime('%d/%m/%Y %H:%M')} UTC")
        doc.add_paragraph(f"Gerente responsável: {gerente}")
        if comentario:
            doc.add_paragraph(f"Observações adicionais: {comentario}")
        cleaned = self._strip_html(relatorio_html)
        if cleaned:
            doc.add_paragraph("Resumo do relatório:")
            doc.add_paragraph(cleaned)

        table = doc.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        headers = ["Código", "Descrição", "Quantidade", "Responsável", "Data da Saída", "Observações"]
        hdr_cells = table.rows[0].cells
        for idx, label in enumerate(headers):
            hdr_cells[idx].text = label
        for entry in entries:
            row_cells = table.add_row().cells
            row_cells[0].text = entry.get("codigo", "-") or "-"
            row_cells[1].text = entry.get("descricao", "-") or "-"
            row_cells[2].text = str(entry.get("quantidade", 0))
            row_cells[3].text = entry.get("responsavel", "-") or "-"
            data_saida = entry.get("saida_data")
            row_cells[4].text = data_saida.strftime('%d/%m/%Y %H:%M') if data_saida else "-"
            row_cells[5].text = entry.get("observacoes", "-") or "-"

        doc.add_paragraph("\nAssinatura digital do gerente: ________________________________")
        return doc

    def _to_dict(self, entry: MaterialInventario) -> dict[str, Any]:
        saida = entry.saida
        responsavel = entry.responsavel_nome or (saida.usuario.nome if saida and saida.usuario else None)
        return {
            "id": entry.id,
            "saida_id": entry.saida_id,
            "data": entry.data_registro,
            "codigo": entry.codigo_item or (saida.codigo_item if saida else ""),
            "descricao": entry.descricao,
            "observacoes": entry.observacoes,
            "quantidade": entry.quantidade,
            "matricula": entry.matricula or (saida.matricula if saida else ""),
            "responsavel": responsavel or "",
            "saida_data": saida.data_saida if saida else None,
        }

    def _map_saida(self, saida: Saida) -> dict[str, Any]:
        return {
            "id": saida.id_saida,
            "codigo": saida.codigo_item,
            "descricao": saida.item.descricao if saida.item else "",
            "quantidade": saida.quantidade,
            "matricula": saida.matricula,
            "usuario": saida.usuario.nome if saida.usuario else "",
            "data": saida.data_saida,
        }


material_inventory_service = MaterialInventoryService()
