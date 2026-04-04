from __future__ import annotations

import csv
from io import BytesIO
from io import StringIO
from typing import Any

from flask import Blueprint, Response, abort, jsonify, render_template, request, send_file, url_for
from flask_login import current_user, login_required
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..services.analytics_service import analytics_service
from ..utils.report_branding import get_company_header_lines


blueprint = Blueprint("analytics", __name__, url_prefix="/analitica")


def _current_user_is_admin() -> bool:
    if not getattr(current_user, "is_authenticated", False):
        return False
    for attr_name in ("is_admin", "admin"):
        raw_value = getattr(current_user, attr_name, False)
        if isinstance(raw_value, bool):
            if raw_value:
                return True
            continue
        normalized = str(raw_value or "").strip().lower()
        if normalized in {"1", "true", "t", "yes", "sim", "admin"}:
            return True
    return False


def _require_admin() -> None:
    if not _current_user_is_admin():
        abort(403)


def _build_export_query_args(payload: dict[str, Any]) -> dict[str, str]:
    filters = payload.get("filters") or {}
    args: dict[str, str] = {}
    for source_key, target_key in (
        ("selected_exercise", "exercicio"),
        ("selected_period_preset", "preset"),
        ("selected_start_date", "data_inicio"),
        ("selected_end_date", "data_fim"),
        ("selected_local", "local"),
        ("selected_category", "categoria"),
        ("selected_employee", "matricula"),
        ("search", "search"),
        ("sort_by", "sort_by"),
        ("sort_dir", "sort_dir"),
    ):
        value = str(filters.get(source_key) or "").strip()
        if value:
            args[target_key] = value
    return args


def _request_dashboard_filters(*, page_default: int = 1, per_page_default: int = 25) -> dict[str, Any]:
    return {
        "exercise_label": (request.args.get("exercicio") or "").strip() or None,
        "period_preset": (request.args.get("preset") or "").strip() or None,
        "start_date_value": (request.args.get("data_inicio") or "").strip() or None,
        "end_date_value": (request.args.get("data_fim") or "").strip() or None,
        "local_name": (request.args.get("local") or "").strip() or None,
        "category_name": (request.args.get("categoria") or "").strip() or None,
        "employee_id": (request.args.get("matricula") or "").strip() or None,
        "search": request.args.get("search", ""),
        "page": request.args.get("page", page_default, type=int),
        "per_page": request.args.get("per_page", per_page_default, type=int),
        "sort_by": (request.args.get("sort_by") or "").strip() or None,
        "sort_dir": (request.args.get("sort_dir") or "").strip() or None,
    }


def _current_query_args() -> dict[str, str]:
    return {key: value for key, value in request.args.to_dict(flat=True).items() if str(value or "").strip()}


def _activity_source_label(value: Any) -> str:
    if str(value or "").strip().lower() == "structured":
        return "Estruturada"
    return "Heuristica"


def _build_workbook(payload: dict[str, Any]) -> BytesIO:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Resumo"

    company_lines = get_company_header_lines()
    current_row = 1
    for line in company_lines[:4]:
        summary_sheet.cell(row=current_row, column=1, value=line)
        summary_sheet.cell(row=current_row, column=1).font = Font(bold=True, size=12)
        current_row += 1

    header = payload.get("header") or {}
    context = payload.get("context") or {}
    kpis = payload.get("kpis") or {}
    filters = payload.get("filters") or {}
    maintenance = payload.get("maintenance") or {}
    charts = payload.get("charts") or {}

    summary_sheet.cell(row=current_row, column=1, value="Central Analitica - Resumo executivo")
    summary_sheet.cell(row=current_row, column=1).font = Font(bold=True, size=14)
    current_row += 1
    summary_sheet.cell(row=current_row, column=1, value=f"Periodo: {header.get('period_label') or 'Atual'}")
    summary_sheet.cell(row=current_row, column=2, value=f"Atualizado em: {header.get('generated_at_label') or '-'}")
    current_row += 2

    summary_rows = [
        ("Escopo", context.get("scope_title") or "Central Analitica"),
        ("Descricao", context.get("scope_subtitle") or "-"),
        ("Movimentacoes", kpis.get("movements") or 0),
        ("Quantidade movimentada", kpis.get("quantity") or 0),
        ("Valor movimentado", kpis.get("value") or 0),
        ("Ticket medio", kpis.get("average_value") or 0),
        ("Quantidade media", kpis.get("average_quantity") or 0),
        ("Colaboradores", kpis.get("collaborators") or 0),
        ("Itens distintos", kpis.get("items") or 0),
        ("Estoque atual", kpis.get("stock_total") or 0),
        ("Estoque em atencao", kpis.get("stock_low") or 0),
        ("Valor estimado de entradas", kpis.get("entry_value") or 0),
        ("Cobertura de atividade derivada", f"{maintenance.get('coverage_percent') or 0}%"),
    ]

    for label, value in summary_rows:
        summary_sheet.cell(row=current_row, column=1, value=label).font = Font(bold=True)
        summary_sheet.cell(row=current_row, column=2, value=value)
        current_row += 1

    current_row += 1
    summary_sheet.cell(row=current_row, column=1, value="Filtros aplicados").font = Font(bold=True, size=12)
    current_row += 1
    filter_rows = [
        ("Exercicio", filters.get("selected_exercise") or "Atual"),
        ("Preset rapido", filters.get("selected_period_preset_label") or "-"),
        ("Data inicial", filters.get("selected_start_date") or "-"),
        ("Data final", filters.get("selected_end_date") or "-"),
        ("Local", filters.get("selected_local") or "Todos"),
        ("Categoria", filters.get("selected_category") or "Todas"),
        ("Funcionario", filters.get("selected_employee") or "Todos"),
        ("Busca", filters.get("search") or "-"),
        ("Ordenacao", f"{filters.get('sort_by') or 'data'} / {filters.get('sort_dir') or 'desc'}"),
    ]
    for label, value in filter_rows:
        summary_sheet.cell(row=current_row, column=1, value=label).font = Font(bold=True)
        summary_sheet.cell(row=current_row, column=2, value=value)
        current_row += 1

    current_row += 1
    summary_sheet.cell(row=current_row, column=1, value="Insights").font = Font(bold=True, size=12)
    current_row += 1
    for insight in payload.get("insights") or []:
        summary_sheet.cell(row=current_row, column=1, value=insight.get("title"))
        summary_sheet.cell(row=current_row, column=2, value=insight.get("value"))
        summary_sheet.cell(row=current_row, column=3, value=insight.get("caption"))
        current_row += 1

    for column, width in {1: 28, 2: 30, 3: 72}.items():
        summary_sheet.column_dimensions[get_column_letter(column)].width = width

    detail_sheet = workbook.create_sheet("Drill-down")

    headers = [
        "Data",
        "Colaborador",
        "Cargo",
        "Item",
        "Codigo",
        "Categoria",
        "Local",
        "Quantidade",
        "Unidade",
        "Valor Total",
        "Tipo",
        "Atividade Derivada",
        "Origem Atividade",
        "Atividade Estruturada",
        "OS",
        "Centro de Custo",
        "Observacao",
    ]

    header_fill = PatternFill(fill_type="solid", fgColor="0F172A")
    header_font = Font(color="FFFFFF", bold=True)
    border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    detail_row = 1
    for column, header in enumerate(headers, start=1):
        cell = detail_sheet.cell(row=detail_row, column=column, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

    for row_data in (payload.get("table") or {}).get("all_rows") or []:
        detail_row += 1
        values = [
            row_data.get("data_saida_label"),
            row_data.get("colaborador_nome"),
            row_data.get("cargo"),
            row_data.get("descricao_item"),
            row_data.get("codigo_item"),
            row_data.get("categoria"),
            row_data.get("local"),
            row_data.get("quantidade_base"),
            row_data.get("unit_base"),
            row_data.get("valor_total"),
            row_data.get("tipo_consumo"),
            row_data.get("atividade_label"),
            _activity_source_label(row_data.get("atividade_source")),
            row_data.get("atividade_estruturada_label"),
            row_data.get("ordem_servico"),
            row_data.get("centro_custo"),
            row_data.get("observacao"),
        ]
        for column, value in enumerate(values, start=1):
            cell = detail_sheet.cell(row=detail_row, column=column, value=value)
            cell.border = border
            if column in {8, 10}:
                cell.alignment = Alignment(horizontal="right")

    widths = {
        1: 18,
        2: 26,
        3: 22,
        4: 34,
        5: 18,
        6: 22,
        7: 20,
        8: 14,
        9: 12,
        10: 16,
        11: 16,
        12: 22,
        13: 18,
        14: 24,
        15: 18,
        16: 22,
        17: 36,
    }
    for column, width in widths.items():
        detail_sheet.column_dimensions[get_column_letter(column)].width = width

    activity_sheet = workbook.create_sheet("Atividades")
    activity_sheet.cell(row=1, column=1, value="Atividade derivada").font = Font(bold=True, size=13)
    activity_sheet.cell(row=2, column=1, value="Cobertura heuristica")
    activity_sheet.cell(row=2, column=2, value=maintenance.get("coverage_percent") or 0)
    activity_sheet.cell(row=3, column=1, value="Ruido Central de Kits")
    activity_sheet.cell(row=3, column=2, value=maintenance.get("noise_percent") or 0)
    activity_sheet.cell(row=4, column=1, value="Sem classificacao")
    activity_sheet.cell(row=4, column=2, value=maintenance.get("unclassified_percent") or 0)

    group_header_row = 6
    for column, header in enumerate(("Grupo", "Movimentos", "Valor", "Quantidade"), start=1):
        cell = activity_sheet.cell(row=group_header_row, column=column, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border

    activity_row = group_header_row
    for group in maintenance.get("groups") or []:
        activity_row += 1
        activity_sheet.cell(row=activity_row, column=1, value=group.get("label")).border = border
        activity_sheet.cell(row=activity_row, column=2, value=group.get("count")).border = border
        activity_sheet.cell(row=activity_row, column=3, value=group.get("value")).border = border
        activity_sheet.cell(row=activity_row, column=4, value=group.get("quantity")).border = border

    activity_row += 2
    for column, header in enumerate(("Frente bruta", "Movimentos", "Valor"), start=1):
        cell = activity_sheet.cell(row=activity_row, column=column, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
    for front in maintenance.get("fronts") or []:
        activity_row += 1
        activity_sheet.cell(row=activity_row, column=1, value=front.get("label")).border = border
        activity_sheet.cell(row=activity_row, column=2, value=front.get("count")).border = border
        activity_sheet.cell(row=activity_row, column=3, value=front.get("value")).border = border

    activity_row += 2
    activity_sheet.cell(row=activity_row, column=1, value="Notas heuristicas").font = Font(bold=True)
    for note in maintenance.get("notes") or []:
        activity_row += 1
        activity_sheet.cell(row=activity_row, column=1, value=note)

    for column, width in {1: 34, 2: 18, 3: 18, 4: 18}.items():
        activity_sheet.column_dimensions[get_column_letter(column)].width = width

    distribution_sheet = workbook.create_sheet("Distribuicao")
    distribution_sheet.cell(row=1, column=1, value="Top locais").font = Font(bold=True, size=12)
    distribution_sheet.cell(row=1, column=4, value="Top categorias").font = Font(bold=True, size=12)
    distribution_sheet.cell(row=1, column=7, value="Top itens").font = Font(bold=True, size=12)

    for index, row in enumerate(charts.get("locations") or [], start=2):
        distribution_sheet.cell(row=index, column=1, value=row.get("local"))
        distribution_sheet.cell(row=index, column=2, value=row.get("total_valor"))
        distribution_sheet.cell(row=index, column=3, value=row.get("saidas"))

    for index, row in enumerate(charts.get("categories") or [], start=2):
        distribution_sheet.cell(row=index, column=4, value=row.get("categoria"))
        distribution_sheet.cell(row=index, column=5, value=row.get("total_valor"))
        distribution_sheet.cell(row=index, column=6, value=row.get("saidas"))

    for index, row in enumerate(payload.get("top_items") or [], start=2):
        distribution_sheet.cell(row=index, column=7, value=row.get("descricao_item"))
        distribution_sheet.cell(row=index, column=8, value=row.get("valor_total"))
        distribution_sheet.cell(row=index, column=9, value=row.get("share_percent"))

    for column, width in {1: 24, 2: 16, 3: 12, 4: 24, 5: 16, 6: 12, 7: 30, 8: 16, 9: 14}.items():
        distribution_sheet.column_dimensions[get_column_letter(column)].width = width

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


@blueprint.get("/")
@login_required
def index():
    _require_admin()
    shell = analytics_service.get_shell_payload(
        **_request_dashboard_filters(),
    )
    export_args = _build_export_query_args(shell)
    pdf_url = url_for("inventory.consumption_report_pdf", **{k: v for k, v in export_args.items() if k != "search"})
    xlsx_url = url_for("analytics.export_xlsx", **export_args)
    csv_url = url_for("analytics.export_csv", **export_args)
    query_args = _current_query_args()
    return render_template(
        "analytics/index.html",
        analytics=shell,
        export_pdf_url=pdf_url,
        export_xlsx_url=xlsx_url,
        export_csv_url=csv_url,
        pdf_is_limited=bool(shell.get("filters", {}).get("selected_start_date") or shell.get("filters", {}).get("selected_end_date") or shell.get("filters", {}).get("search")),
        api_urls={
            "overview": url_for("analytics.api_overview", **query_args),
            "charts": url_for("analytics.api_charts", **query_args),
            "table": url_for("analytics.api_table", **query_args),
            "maintenance": url_for("analytics.api_maintenance", **query_args),
        },
    )


@blueprint.get("/api/overview")
@login_required
def api_overview():
    _require_admin()
    return jsonify(analytics_service.get_overview_widget(**_request_dashboard_filters()))


@blueprint.get("/api/charts")
@login_required
def api_charts():
    _require_admin()
    return jsonify(analytics_service.get_charts_widget(**_request_dashboard_filters()))


@blueprint.get("/api/table")
@login_required
def api_table():
    _require_admin()
    return jsonify(analytics_service.get_table_widget(**_request_dashboard_filters()))


@blueprint.get("/api/maintenance")
@login_required
def api_maintenance():
    _require_admin()
    return jsonify(analytics_service.get_maintenance_widget(**_request_dashboard_filters()))


@blueprint.get("/export.xlsx")
@login_required
def export_xlsx():
    _require_admin()
    export_filters = _request_dashboard_filters(page_default=1, per_page_default=100000)
    export_filters["page"] = 1
    export_filters["per_page"] = 100000
    payload = analytics_service.get_dashboard_payload(**export_filters)
    workbook = _build_workbook(payload)
    period_label = ((payload.get("header") or {}).get("period_label") or "atual").replace("/", "_").replace(" ", "_")
    filename = f"central_analitica_{period_label}.xlsx"
    return send_file(
        workbook,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


@blueprint.get("/export.csv")
@login_required
def export_csv():
    _require_admin()
    export_filters = _request_dashboard_filters(page_default=1, per_page_default=100000)
    export_filters["page"] = 1
    export_filters["per_page"] = 100000
    payload = analytics_service.get_dashboard_payload(**export_filters)
    output = StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow([
        "Data",
        "Colaborador",
        "Cargo",
        "Item",
        "Codigo",
        "Categoria",
        "Local",
        "Quantidade",
        "Unidade",
        "Valor Total",
        "Tipo",
        "Atividade Derivada",
        "Origem Atividade",
        "Atividade Estruturada",
        "OS",
        "Centro de Custo",
        "Observacao",
    ])
    for row in (payload.get("table") or {}).get("all_rows") or []:
        writer.writerow([
            row.get("data_saida_label"),
            row.get("colaborador_nome"),
            row.get("cargo"),
            row.get("descricao_item"),
            row.get("codigo_item"),
            row.get("categoria"),
            row.get("local"),
            row.get("quantidade_base"),
            row.get("unit_base"),
            row.get("valor_total"),
            row.get("tipo_consumo"),
            row.get("atividade_label"),
            _activity_source_label(row.get("atividade_source")),
            row.get("atividade_estruturada_label"),
            row.get("ordem_servico"),
            row.get("centro_custo"),
            row.get("observacao"),
        ])
    period_label = ((payload.get("header") or {}).get("period_label") or "atual").replace("/", "_").replace(" ", "_")
    csv_content = "\ufeff" + output.getvalue()
    return Response(
        csv_content,
        mimetype="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": f"attachment; filename=central_analitica_{period_label}.csv",
        },
    )