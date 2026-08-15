"""View para controle de custódia de ferramentas."""
from __future__ import annotations

from flask import Blueprint, current_app, render_template, request, flash, redirect, url_for, jsonify, send_file
from flask_login import login_required, current_user
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy import or_

from ..extensions import db
from ..models import RetiradaFerramenta, Saida, Usuario
from ..services.ferramentas import ferramentas_service
from ..services.tool_custody_service import tool_custody_service
from ..utils.time_service import TimeService

bp = Blueprint("tool_custody", __name__, url_prefix="/controle-ferramentas")


@bp.route("/")
@login_required
def index():
    """Lista todos os funcionários com ferramentas ativas."""
    # Buscar parâmetros de pesquisa e filtro
    search = request.args.get("search", "").strip()
    setor_filter = request.args.get("setor", "")
    alert_only = request.args.get("alert_only", "") == "true"
    tipo_custodia_filter = request.args.get("tipo_custodia", "todos")  # Padrão: mostrar todos
    
    # Obter lista de funcionários
    employees = tool_custody_service.get_all_employees_with_tools()
    
    # Aplicar filtros
    if search:
        search_lower = search.lower()
        employees = [
            e for e in employees
            if search_lower in e["nome"].lower() 
            or search_lower in e["matricula"].lower()
            or any(search_lower in tool.get("codigo_item", "").lower() for tool in e.get("tools", []))
        ]
    
    if setor_filter:
        employees = [e for e in employees if e["setor"] == setor_filter]
    
    if alert_only:
        employees = [e for e in employees if e["has_alerts"]]
    
    # Filtrar por tipo de custódia (diária/permanente/todos)
    if tipo_custodia_filter == "diaria":
        employees = [e for e in employees if e.get("has_temporaria", False)]
    elif tipo_custodia_filter == "permanente":
        employees = [e for e in employees if e.get("has_permanente", False)]
    # Se "todos", não filtra
    
    # Obter lista de setores únicos para o filtro
    all_employees = tool_custody_service.get_all_employees_with_tools()
    setores = sorted(list(set(e["setor"] for e in all_employees)))
    
    # Estatísticas
    stats = tool_custody_service.get_statistics()
    lost_broken_tools = tool_custody_service.list_lost_broken_tools(limit=500)
    lost_broken_stats = {
        "total": len(lost_broken_tools),
        "perdidas": sum(1 for row in lost_broken_tools if row.get("kind") == "perdida"),
        "quebradas": sum(1 for row in lost_broken_tools if row.get("kind") == "quebrada"),
        "valor_total": sum(float(row.get("valor_total") or 0.0) for row in lost_broken_tools),
    }
    for row in lost_broken_tools:
        foto_path = row.get("foto_path")
        row["foto_url"] = url_for("static", filename=foto_path) if foto_path else None
    
    return render_template(
        "tool_custody/index.html",
        employees=employees,
        setores=setores,
        search=search,
        setor_filter=setor_filter,
        alert_only=alert_only,
        tipo_custodia_filter=tipo_custodia_filter,
        stats=stats,
        lost_broken_tools=lost_broken_tools,
        lost_broken_stats=lost_broken_stats,
    )


@bp.route("/funcionario/<matricula>")
@login_required
def employee_detail(matricula: str):
    """Exibe detalhes de um funcionário e suas ferramentas."""
    from ..models import TelegramUser, Usuario as UsuarioModel
    
    employee = tool_custody_service.get_employee_details(matricula)
    
    if not employee:
        flash("Funcionário não encontrado", "danger")
        return redirect(url_for("tool_custody.index"))
    
    # Buscar todos os usuários Telegram habilitados para mostrar na UI
    telegram_users = (
        db.session.query(TelegramUser, UsuarioModel)
        .join(UsuarioModel, TelegramUser.matricula == UsuarioModel.matricula)
        .filter(TelegramUser.enabled == True)
        .all()
    )
    recipients = [user.nome for _, user in telegram_users] if telegram_users else []
    
    return render_template(
        "tool_custody/detail.html", 
        employee=employee,
        telegram_recipients=recipients
    )


@bp.route("/funcionario/<matricula>/foto", methods=["POST"])
@login_required
def upload_employee_photo(matricula: str):
    """Faz upload da foto do funcionário usada na tela de custódia."""
    file = request.files.get("photo")

    if not file or not file.filename:
        flash("Selecione uma imagem para enviar", "warning")
        return redirect(url_for("tool_custody.employee_detail", matricula=matricula))

    try:
        tool_custody_service.upload_employee_photo(file, matricula)
        flash("Foto do funcionário atualizada com sucesso!", "success")
    except ValueError as e:
        flash(str(e), "danger")
    except Exception as e:
        flash(f"Erro ao atualizar foto: {str(e)}", "danger")

    return redirect(url_for("tool_custody.employee_detail", matricula=matricula))


@bp.route("/devolucao/<int:saida_id>", methods=["POST"])
@login_required
def return_tool(saida_id: int):
    """Registra devolução de ferramenta."""
    try:
        observacao = request.form.get("observacao", "").strip()
        source = (request.form.get("source") or "saida").strip().lower()
        if source == "retirada_ferramenta":
            ferramentas_service.devolver_ferramenta(saida_id, observacao or None)
        else:
            tool_custody_service.register_return(saida_id, observacao or None)
        flash("Ferramenta devolvida com sucesso!", "success")
    except ValueError as e:
        flash(str(e), "danger")
    except SQLAlchemyError as e:
        flash(f"Erro ao registrar devolução: {str(e)}", "danger")
    
    # Redirecionar de volta para a página do funcionário
    matricula = request.form.get("matricula")
    if matricula:
        return redirect(url_for("tool_custody.employee_detail", matricula=matricula))
    return redirect(url_for("tool_custody.index"))


@bp.route("/devolucao/<int:saida_id>/api", methods=["POST"])
@login_required
def return_tool_api(saida_id: int):
    """Registra devolucao de ferramenta via API (sem redirect)."""
    try:
        observacao = request.form.get("observacao", "").strip()
        source = (request.form.get("source") or "saida").strip().lower()
        if source == "retirada_ferramenta":
            ferramentas_service.devolver_ferramenta(saida_id, observacao or None)
        else:
            tool_custody_service.register_return(saida_id, observacao or None)
        return jsonify({"success": True, "message": "Ferramenta devolvida com sucesso"})
    except ValueError as e:
        message = str(e)

        # Idempotência: se já foi devolvida, consideramos sucesso (o card deve sumir do feed).
        if message.strip().lower() == "ferramenta já foi devolvida":
            return jsonify({"success": True, "message": message})

        # Fallback: se o id não resolve (dados legados/ids divergentes), tenta por matrícula+código.
        if message.strip().lower() == "retirada não encontrada":
            from ..models import Item, Saida

            matricula = (request.form.get("matricula") or "").strip()
            codigo_item = (request.form.get("codigo_item") or "").strip()

            if matricula and codigo_item:
                saida = (
                    db.session.query(Saida)
                    .join(Item, Saida.codigo_item == Item.codigo_item)
                    .filter(
                        Saida.matricula == matricula,
                        Saida.codigo_item == codigo_item,
                        Item.categoria == "Ferramentas",
                    )
                    .order_by(Saida.data_saida.desc())
                    .first()
                )

                if saida:
                    try:
                        tool_custody_service.register_return(saida.id_saida, observacao or None)
                        return jsonify({"success": True, "message": "Ferramenta devolvida com sucesso"})
                    except ValueError as inner:
                        inner_msg = str(inner)
                        if inner_msg.strip().lower() == "ferramenta já foi devolvida":
                            return jsonify({"success": True, "message": inner_msg})
                        return jsonify({"success": False, "message": inner_msg}), 400

        return jsonify({"success": False, "message": message}), 400
    except SQLAlchemyError as e:
        return jsonify({"success": False, "message": f"Erro ao registrar devolucao: {str(e)}"}), 500


@bp.route("/quebra/<int:saida_id>/relatorio.pdf")
@login_required
def damage_tool_report(saida_id: int):
    """Gera PDF de quebra/perda para assinatura do funcionário."""
    from io import BytesIO

    source = (request.args.get("source") or "saida").strip().lower()
    kind = (request.args.get("kind") or "quebrada").strip().lower()
    source_norm = "retirada_ferramenta" if source == "retirada_ferramenta" else "saida"
    kind_label = "Perdida" if kind in {"perdida", "perda", "lost"} else "Quebrada"

    saida = Saida.query.get(saida_id) if source_norm == "saida" else None
    retirada = db.session.get(RetiradaFerramenta, saida_id) if source_norm == "retirada_ferramenta" else None

    record = saida or retirada
    if record is None:
        flash("Registro de ferramenta não encontrado para gerar o relatório.", "danger")
        return redirect(url_for("tool_custody.index"))

    item = getattr(record, "item", None) or None
    codigo_item = getattr(record, "codigo_item", None) or "N/D"
    descricao = getattr(item, "descricao", None) or "Ferramenta"
    matricula = getattr(record, "matricula", None) or ""
    usuario = db.session.query(Usuario).filter_by(matricula=matricula).first() if matricula else None
    local_servico = getattr(record, "local_servico", None) or "Não informado"
    observacao = getattr(record, "observacao", None) or ""
    data_retirada = getattr(record, "data_retirada", None) or getattr(record, "data_saida", None)
    data_text = TimeService.format_local(data_retirada, "%d/%m/%Y %H:%M") if data_retirada else "N/D"

    html = f"""
    <!doctype html>
    <html lang="pt-BR">
    <head>
      <meta charset="utf-8">
      <title>Relatório de ferramenta {kind_label.lower()}</title>
      <style>
        body {{ font-family: Arial, sans-serif; margin: 32px; color: #111827; }}
        h1 {{ font-size: 24px; margin-bottom: 10px; }}
        .card {{ border: 1px solid #d1d5db; border-radius: 8px; padding: 18px; margin-top: 16px; }}
        .grid {{ display: grid; grid-template-columns: 1fr 1fr; gap: 12px; margin-top: 14px; }}
        .label {{ font-size: 12px; color: #6b7280; text-transform: uppercase; margin-bottom: 4px; }}
        .value {{ font-size: 15px; font-weight: 600; }}
        .signature {{ border-top: 2px solid #111827; margin-top: 60px; padding-top: 10px; width: 300px; }}
      </style>
    </head>
    <body>
      <h1>Relatório de ferramenta {kind_label}</h1>
      <p>Este documento registra a ocorrência de ferramenta {kind_label.lower()} e deve ser assinado pelo colaborador responsável.</p>
      <div class="card">
        <div class="grid">
          <div><div class="label">Colaborador</div><div class="value">{(usuario.nome if usuario else matricula) or 'N/D'}</div></div>
          <div><div class="label">Matrícula</div><div class="value">{matricula or 'N/D'}</div></div>
          <div><div class="label">Ferramenta</div><div class="value">{descricao}</div></div>
          <div><div class="label">Código</div><div class="value">{codigo_item}</div></div>
          <div><div class="label">Data retirada</div><div class="value">{data_text}</div></div>
          <div><div class="label">Local</div><div class="value">{local_servico}</div></div>
        </div>
        <div style="margin-top: 18px;"><div class="label">Motivo / observação</div><div class="value">{observacao or 'Motivo não informado'}</div></div>
      </div>
      <div class="signature">
        <div>Assinatura do funcionário:</div>
        <div style="height: 32px;"></div>
      </div>
    </body>
    </html>
    """

    try:
        from ..utils.html_pdf import render_html_to_pdf

        result = render_html_to_pdf(html=html, base_url=current_app.root_path)
        return send_file(
            BytesIO(result.pdf_bytes),
            mimetype="application/pdf",
            as_attachment=True,
            download_name=f"relatorio_ferramenta_{kind_label.lower()}_{saida_id}.pdf",
        )
    except Exception as exc:
        flash(f"Não foi possível gerar o relatório PDF: {exc}", "danger")
        return redirect(url_for("tool_custody.index"))


@bp.route("/quebra/<int:saida_id>", methods=["POST"])
@login_required
def damage_tool(saida_id: int):
    """Registra ferramenta quebrada/danificada."""
    try:
        observacao = request.form.get("observacao", "").strip()
        source = (request.form.get("source") or "saida").strip().lower()
        report_signed = request.form.get("report_signed") == "1"
        if not report_signed:
            flash("Primeiro gere o relatório em PDF e confirme que o funcionário assinou antes de registrar a perda/quebra.", "warning")
            return redirect(request.referrer or url_for("tool_custody.index"))
        tool_custody_service.register_damage(saida_id, observacao or None, source=source)
        flash("Quebra/dano registrado com sucesso! O saldo foi ajustado no sistema.", "warning")
    except ValueError as e:
        flash(str(e), "danger")
    except SQLAlchemyError as e:
        flash(f"Erro ao registrar dano: {str(e)}", "danger")
    
    matricula = request.form.get("matricula")
    if matricula:
        return redirect(url_for("tool_custody.employee_detail", matricula=matricula))
    return redirect(url_for("tool_custody.index"))


@bp.route("/quebra/<int:saida_id>/api", methods=["POST"])
@login_required
def damage_tool_api(saida_id: int):
    """Registra ferramenta quebrada/perdida via API."""
    try:
        source = (request.form.get("source") or "saida").strip()
        kind = (request.form.get("kind") or request.form.get("tipo") or "quebrada").strip()
        observacao = request.form.get("observacao", "").strip()
        report_signed = request.form.get("report_signed") == "1"
        if not report_signed:
            return jsonify({"success": False, "message": "É necessário gerar e assinar o relatório PDF antes de concluir a baixa."}), 400
        tool_custody_service.register_damage(saida_id, observacao or None, kind=kind, source=source)
        return jsonify({"success": True, "message": "Ferramenta registrada como quebrada/perdida"})
    except ValueError as e:
        return jsonify({"success": False, "message": str(e)}), 400
    except SQLAlchemyError as e:
        return jsonify({"success": False, "message": f"Erro ao registrar ocorrência: {str(e)}"}), 500


@bp.route("/reparo/<int:saida_id>", methods=["POST"])
@login_required
def repair_tool(saida_id: int):
    """Registra envio para reparo."""
    try:
        observacao = request.form.get("observacao", "").strip()
        tool_custody_service.register_repair(saida_id, observacao or None)
        flash("Envio para reparo registrado com sucesso!", "info")
    except ValueError as e:
        flash(str(e), "danger")
    except SQLAlchemyError as e:
        flash(f"Erro ao registrar reparo: {str(e)}", "danger")
    
    matricula = request.form.get("matricula")
    if matricula:
        return redirect(url_for("tool_custody.employee_detail", matricula=matricula))
    return redirect(url_for("tool_custody.index"))


@bp.route("/funcionario/<matricula>/adicionar-ferramenta", methods=["POST"])
@login_required
def add_tool_to_employee(matricula: str):
    """Adiciona uma ferramenta diretamente à custódia do funcionário."""
    try:
        codigo_item = (request.form.get("codigo_item") or request.form.get("codigo") or "").strip()
        quantidade = int((request.form.get("quantidade") or "1").strip())
        tipo_custodia = (request.form.get("tipo_custodia") or "temporaria").strip()
        local_servico = (request.form.get("local_servico") or "").strip() or None
        observacao = (request.form.get("observacao") or "").strip() or None

        saida_id = tool_custody_service.add_tool_to_employee(
            matricula=matricula,
            codigo_item=codigo_item,
            quantidade=quantidade,
            tipo_custodia=tipo_custodia,
            local_servico=local_servico,
            observacao=observacao,
        )

        tipo_label = "permanente" if str(tipo_custodia or "").strip().lower() == "permanente" else "diária"
        flash(f"Ferramenta adicionada à custódia {tipo_label}. Saída #{saida_id} registrada com sucesso!", "success")
    except ValueError as e:
        flash(str(e), "danger")
    except SQLAlchemyError as e:
        flash(f"Erro ao adicionar ferramenta à custódia: {str(e)}", "danger")

    return redirect(url_for("tool_custody.employee_detail", matricula=matricula))


@bp.route("/transferir/<int:saida_id>", methods=["POST"])
@login_required
def transfer_tool(saida_id: int):
    """Transfere a custódia de uma ferramenta para outro funcionário."""
    matricula_origem = (request.form.get("matricula") or "").strip()
    try:
        nova_matricula = (request.form.get("nova_matricula") or "").strip()
        local_servico = (request.form.get("local_servico") or "").strip() or None
        observacao = (request.form.get("observacao") or "").strip() or None

        novo_saida_id = tool_custody_service.transfer_tool(
            saida_id=saida_id,
            nova_matricula=nova_matricula,
            local_servico=local_servico,
            observacao=observacao,
        )
        flash(f"Custódia transferida com sucesso! Nova saída #{novo_saida_id} registrada.", "success")
    except ValueError as e:
        flash(str(e), "danger")
    except SQLAlchemyError as e:
        flash(f"Erro ao transferir ferramenta: {str(e)}", "danger")

    if matricula_origem:
        return redirect(url_for("tool_custody.employee_detail", matricula=matricula_origem))
    return redirect(url_for("tool_custody.index"))


@bp.route("/api/stats")
@login_required
def api_stats():
    """API endpoint para estatísticas atualizadas."""
    stats = tool_custody_service.get_statistics()
    return jsonify(stats)


@bp.route("/custodia/<int:saida_id>", methods=["POST"])
@login_required
def toggle_custody_type(saida_id: int):
    """Alterna tipo de custódia entre permanente e temporária."""
    try:
        novo_tipo = (request.form.get("novo_tipo", "temporaria") or "").strip().lower()
        # Compat: algumas telas/instalações antigas usavam "diaria" para empréstimo temporário.
        if novo_tipo in {"diaria", "diária", "daily", "d"}:
            novo_tipo = "temporaria"

        if novo_tipo not in ["temporaria", "permanente"]:
            raise ValueError("Tipo de custódia inválido")
        
        tool_custody_service.change_custody_type(saida_id, novo_tipo)
        
        tipo_label = "permanente" if novo_tipo == "permanente" else "temporária"
        flash(f"Tipo de custódia alterado para: {tipo_label}", "success")
    except ValueError as e:
        flash(str(e), "danger")
    except SQLAlchemyError as e:
        flash(f"Erro ao alterar tipo de custódia: {str(e)}", "danger")
    
    matricula = request.form.get("matricula")
    if matricula:
        return redirect(url_for("tool_custody.employee_detail", matricula=matricula))
    return redirect(url_for("tool_custody.index"))


@bp.route("/alertar/<int:saida_id>", methods=["POST"])
@login_required
def send_telegram_alert(saida_id: int):
    """Envia alerta via Telegram sobre ferramenta não devolvida."""
    try:
        result = tool_custody_service.send_telegram_alert(saida_id)
        if result:
            flash("Alerta enviado via Telegram com sucesso!", "success")
        else:
            flash("Erro ao enviar alerta via Telegram", "warning")
    except ValueError as e:
        flash(str(e), "danger")
    except Exception as e:
        flash(f"Erro ao enviar alerta: {str(e)}", "danger")
    
    matricula = request.form.get("matricula")
    if matricula:
        return redirect(url_for("tool_custody.employee_detail", matricula=matricula))
    return redirect(url_for("tool_custody.index"))


@bp.route("/alertar-funcionario/<matricula>", methods=["POST"])
@login_required
def send_grouped_alert(matricula: str):
    """Envia alerta agrupado de todas as ferramentas atrasadas do funcionário."""
    try:
        result = tool_custody_service.send_grouped_alert_by_employee(matricula)
        if result:
            flash("Alerta agrupado enviado via Telegram com sucesso!", "success")
        else:
            flash("Erro ao enviar alerta via Telegram", "warning")
    except ValueError as e:
        flash(str(e), "danger")
    except Exception as e:
        flash(f"Erro ao enviar alerta: {str(e)}", "danger")
    
    return redirect(url_for("tool_custody.employee_detail", matricula=matricula))


@bp.route("/excluir/<int:saida_id>", methods=["POST"])
@login_required
def delete_tool_record(saida_id: int):
    """Exclui registro de ferramenta (APENAS ADMINISTRADORES)."""
    # Verificar se é administrador
    if not current_user.is_authenticated or not current_user.is_admin:
        flash("Acesso negado. Apenas administradores podem excluir registros.", "danger")
        return redirect(url_for("tool_custody.index"))
    
    try:
        tool_custody_service.delete_tool_record(saida_id)
        flash("Registro de ferramenta excluído com sucesso!", "success")
    except ValueError as e:
        flash(str(e), "danger")
    except SQLAlchemyError as e:
        flash(f"Erro ao excluir registro: {str(e)}", "danger")
    
    matricula = request.form.get("matricula")
    if matricula:
        return redirect(url_for("tool_custody.employee_detail", matricula=matricula))
    return redirect(url_for("tool_custody.index"))


@bp.route("/relatorios")
@login_required
def tool_reports():
    """Página de relatórios de ferramentas."""
    from ..utils.time_service import TimeService

    return render_template(
        "tool_custody/reports.html",
        today=TimeService.now_local().date().isoformat(),
    )


@bp.route("/relatorios/gerar", methods=["POST"])
@login_required
def generate_tool_report():
    """Gera relatório de movimentação de ferramentas."""
    from flask import current_app
    from pathlib import Path
    from datetime import datetime, timedelta, timezone
    from ..models import Item, Saida, Usuario, InventarioEvento
    from ..extensions import db
    from ..utils.time_service import TimeService, get_local_tz

    logger = current_app.logger

    def _redirect_back():
        ref = (request.referrer or "").strip()
        # fallback seguro
        fallback = url_for("tool_custody.index")
        if not ref:
            return redirect(fallback)
        # só permitir retornar para páginas internas do módulo
        if "/controle-ferramentas" not in ref:
            return redirect(fallback)
        return redirect(ref)
    
    # Obter parâmetros
    date_from_str = request.form.get("date_from", "").strip()
    date_to_str = request.form.get("date_to", "").strip()
    requested_format = request.form.get("format", "pdf").strip().lower()
    tipo_custodia_filter = request.form.get("tipo_custodia_filter", "todos").strip().lower()
    
    # Validar formato
    if requested_format not in ["pdf", "xlsx"]:
        flash("Formato inválido. Use PDF.", "danger")
        return _redirect_back()

    format_type = "pdf"
    
    # Validar filtro de custódia
    # Compat: instalações antigas usam "diaria" para empréstimo temporário.
    if tipo_custodia_filter in ("diaria",):
        tipo_custodia_filter = "temporaria"
    if tipo_custodia_filter not in ["todos", "permanente", "temporaria"]:
        tipo_custodia_filter = "todos"
    
    # Converter datas
    try:
        if date_from_str:
            date_from = datetime.strptime(date_from_str, "%Y-%m-%d")
        else:
            date_from = datetime.utcnow() - timedelta(days=30)
        
        if date_to_str:
            date_to = datetime.strptime(date_to_str, "%Y-%m-%d")
            date_to = date_to.replace(hour=23, minute=59, second=59)
        else:
            date_to = datetime.utcnow()
    except ValueError:
        flash("Formato de data inválido.", "danger")
        return _redirect_back()
    
    if date_from > date_to:
        flash("Data início deve ser anterior à data fim.", "danger")
        return _redirect_back()

    # As datas do formulário são locais; o banco usa UTC sem tzinfo (padrão do projeto).
    # Convertemos o intervalo local -> UTC para filtrar corretamente.
    try:
        local_tz = get_local_tz()
        date_from_utc = date_from.replace(tzinfo=local_tz).astimezone(timezone.utc).replace(tzinfo=None)
        date_to_utc = date_to.replace(tzinfo=local_tz).astimezone(timezone.utc).replace(tzinfo=None)
    except Exception:
        # fallback: manter comportamento antigo
        date_from_utc = date_from
        date_to_utc = date_to
    
    # Buscar movimentações de ferramentas
    try:
        query = db.session.query(
            Saida.id_saida,
            Saida.quantidade,
            Saida.data_saida,
            Saida.observacao,
            Saida.local_servico,
            Saida.codigo_item,
            Saida.matricula,
            Saida.tipo_custodia,
            Item.descricao.label("item_descricao"),
            Item.marca.label("item_marca"),
            Usuario.nome.label("usuario_nome"),
        ).join(
            Item, Saida.codigo_item == Item.codigo_item, isouter=True
        ).join(
            Usuario, Saida.matricula == Usuario.matricula, isouter=True
        ).filter(
            Saida.data_saida >= date_from_utc,
            Saida.data_saida <= date_to_utc,
            # evitar "sumir" com registros quando o Item foi removido/mudou categoria;
            # aceitamos tanto por categoria quanto por marcação de custódia.
            or_(
                Item.categoria.ilike("Ferrament%"),
                Saida.tipo_custodia.isnot(None),
            ),
        )
        
        # Aplicar filtro de tipo de custódia
        if tipo_custodia_filter != "todos":
            if tipo_custodia_filter == "temporaria":
                query = query.filter(Saida.tipo_custodia.in_(["temporaria", "diaria"]))
            else:
                query = query.filter(Saida.tipo_custodia == tipo_custodia_filter)
        
        query = query.order_by(Saida.data_saida.desc())
        
        results = query.all()

        # Cabeçalho padronizado (pedido do cliente)
        titulo_relatorio = "Relatório de Ferramentas"
        if tipo_custodia_filter == "permanente":
            subtitulo_relatorio = "Custódia Permanente - Uso integral"
        elif tipo_custodia_filter == "temporaria":
            subtitulo_relatorio = "Custódia Diária"
        else:
            subtitulo_relatorio = "Custódia"  # fallback simples
        
        if not results:
            flash("Nenhuma movimentação de ferramenta encontrada no período.", "warning")
            return _redirect_back()
        
        # Buscar devoluções
        devolucoes_query = db.session.query(
            InventarioEvento.codigo_item,
            InventarioEvento.matricula,
            InventarioEvento.quantidade,
            InventarioEvento.data_evento,
            InventarioEvento.descricao,
        ).filter(
            InventarioEvento.tipo == "devolucao",
            InventarioEvento.data_evento >= date_from_utc,
            InventarioEvento.data_evento <= date_to_utc,
        )

        # Robustez: evitar mapear devoluções de materiais com o mesmo código.
        # Se a lista de itens no relatório não for enorme, filtramos no SQL.
        try:
            codes = sorted({r.codigo_item for r in results if getattr(r, "codigo_item", None)})
            if 0 < len(codes) <= 500:
                devolucoes_query = devolucoes_query.filter(InventarioEvento.codigo_item.in_(codes))
        except Exception:
            pass

        devolucoes_query = devolucoes_query.filter(InventarioEvento.matricula.isnot(None)).all()
        
        devolucoes_map = {}
        for dev in devolucoes_query:
            key = (dev.codigo_item, dev.matricula)
            if key not in devolucoes_map:
                devolucoes_map[key] = []
            devolucoes_map[key].append({
                "quantidade": dev.quantidade,
                "data": dev.data_evento,
                "descricao": dev.descricao
            })
        
        # Gerar arquivo
        reports_dir = Path(current_app.instance_path) / "reports"
        reports_dir.mkdir(parents=True, exist_ok=True)
        
        timestamp = TimeService.now_local().strftime('%Y%m%d_%H%M%S')
        period_label = f"{date_from.strftime('%Y%m%d')}_{date_to.strftime('%Y%m%d')}"
        
        if format_type == "xlsx":
            try:
                from openpyxl import Workbook
                from openpyxl.styles import Font, PatternFill, Alignment
            except Exception as e:
                logger.exception("Dependência openpyxl ausente/erro ao importar")
                flash("Não foi possível gerar XLSX (dependência openpyxl).", "danger")
                return _redirect_back()
            
            wb = Workbook()
            ws = wb.active
            ws.title = "Ferramentas"

            # Cabeçalho (título + subtítulo + empresa)
            from ..utils.report_branding import get_company_header_lines

            ws.append([titulo_relatorio])
            ws.append([subtitulo_relatorio])
            for line in get_company_header_lines():
                ws.append([line])
            ws.append([f"Período: {date_from.strftime('%d/%m/%Y')} a {date_to.strftime('%d/%m/%Y')}"])
            ws.append([f"Total de registros: {len(results)}"])
            ws.append([])
            
            # Estilos
            header_fill = PatternFill(start_color="1f2937", end_color="1f2937", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True)
            
            # Cabeçalhos
            headers = ["Data Retirada", "Hora", "Funcionário", "Matrícula", "Ferramenta", "Marca", "Qtd.", "Local", "Devolvido?", "Data Devolução"]
            # A linha do header depende do tamanho do cabeçalho (empresa pode ter 2+ linhas)
            header_row_index = ws.max_row + 1
            ws.append(headers)

            for cell in ws[header_row_index]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")

            # Dados
            for row in results:
                key = (row.codigo_item, row.matricula)
                devolvido = "Não"
                data_devolucao = ""
                
                if key in devolucoes_map and devolucoes_map[key]:
                    devolvido = "Sim"
                    data_devolucao = TimeService.format_local(devolucoes_map[key][0]["data"], "%d/%m/%Y %H:%M")
                
                # Mesclar local e observação
                local_info = row.local_servico or "N/D"
                if row.observacao and row.observacao.strip():
                    local_info = f"{local_info} | {row.observacao[:100]}"
                
                ws.append([
                    TimeService.format_local(row.data_saida, "%d/%m/%Y"),
                    TimeService.format_local(row.data_saida, "%H:%M"),
                    row.usuario_nome or "N/D",
                    row.matricula or "N/D",
                    row.item_descricao or "Ferramenta removida",
                    row.item_marca or "N/D",
                    row.quantidade,
                    local_info[:200],
                    devolvido,
                    data_devolucao,
                ])
            
            # Ajustar larguras
            ws.column_dimensions["A"].width = 12
            ws.column_dimensions["B"].width = 8
            ws.column_dimensions["C"].width = 30
            ws.column_dimensions["D"].width = 12
            ws.column_dimensions["E"].width = 38
            ws.column_dimensions["F"].width = 18
            ws.column_dimensions["G"].width = 8
            ws.column_dimensions["H"].width = 60
            ws.column_dimensions["I"].width = 12
            ws.column_dimensions["J"].width = 20
            
            filename = f"relatorio_ferramentas_{period_label}_{timestamp}.xlsx"
            filepath = reports_dir / filename
            wb.save(str(filepath))
            
            return send_file(
                str(filepath),
                as_attachment=True,
                download_name=filename,
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        
        elif format_type == "pdf":
            # Melhor opção agora: HTML/CSS -> PDF via WeasyPrint (com fallback para ReportLab)
            from ..utils.report_branding import get_company_header_html
            from ..utils.html_pdf import render_html_to_pdf
            from flask import current_app

            period_label_display = f"{date_from.strftime('%d/%m/%Y')} a {date_to.strftime('%d/%m/%Y')}"

            rows = []
            for row in results:
                key = (row.codigo_item, row.matricula)
                devolvido = "Não"
                data_devolucao = ""
                if key in devolucoes_map and devolucoes_map[key]:
                    devolvido = "Sim"
                    data_devolucao = TimeService.format_local(devolucoes_map[key][0]["data"], "%d/%m/%Y")

                local_info = row.local_servico or "N/D"
                if row.observacao and row.observacao.strip():
                    local_info = f"{local_info} | {row.observacao}"

                rows.append(
                    {
                        "data": TimeService.format_local(row.data_saida, "%d/%m/%Y"),
                        "hora": TimeService.format_local(row.data_saida, "%H:%M"),
                        "funcionario": (row.usuario_nome or "N/D"),
                        "ferramenta": (row.item_descricao or "Ferramenta removida"),
                        "qtd": row.quantidade,
                        "local": local_info[:200],
                        "devolvido": devolvido,
                        "data_devolucao": data_devolucao,
                    }
                )

            pdf_filename = f"relatorio_ferramentas_{period_label}_{timestamp}.pdf"

            try:
                from io import BytesIO

                html = render_template(
                    "reports/tool_custody_pdf.html",
                    titulo=titulo_relatorio,
                    subtitulo=subtitulo_relatorio,
                    company_html=get_company_header_html(),
                    periodo=period_label_display,
                    total_registros=len(results),
                    rows=rows,
                )

                result = render_html_to_pdf(html=html, base_url=current_app.root_path)
                buffer = BytesIO(result.pdf_bytes)
                buffer.seek(0)
                return send_file(
                    buffer,
                    as_attachment=True,
                    download_name=pdf_filename,
                    mimetype="application/pdf",
                )
            except Exception:
                logger.exception("Falha ao gerar PDF via WeasyPrint; tentando fallback ReportLab")

            # Fallback: ReportLab (mantém comportamento anterior)
            try:
                from io import BytesIO
                from reportlab.lib import colors
                from reportlab.lib.pagesizes import A4, landscape
                from reportlab.lib.styles import getSampleStyleSheet
                from reportlab.lib.units import cm
                from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
            except Exception:
                flash(
                    "Não foi possível gerar PDF (WeasyPrint e ReportLab indisponíveis).",
                    "danger",
                )
                return _redirect_back()

            buffer = BytesIO()
            doc = SimpleDocTemplate(
                buffer,
                pagesize=landscape(A4),
                leftMargin=0.5 * cm,
                rightMargin=0.5 * cm,
                topMargin=0.5 * cm,
                bottomMargin=0.5 * cm,
                title="Relatório de Ferramentas",
                author="GALINT",
            )

            styles = getSampleStyleSheet()
            body_style = styles["BodyText"]
            body_style.fontSize = 7
            body_style.leading = 8
            story = []

            title_style = styles["Title"]
            title_style.alignment = 1
            title_style.fontSize = 16

            subtitle_style = styles["Normal"]
            subtitle_style.alignment = 1

            story.append(Paragraph(titulo_relatorio, title_style))
            story.append(Paragraph(f"<b>{subtitulo_relatorio}</b>", subtitle_style))
            story.append(Paragraph(get_company_header_html(), subtitle_style))
            story.append(Spacer(1, 0.2 * cm))

            story.append(Paragraph(f"Período: <b>{period_label_display}</b>", styles["Normal"]))
            story.append(Paragraph(f"Total de registros: {len(results)}", styles["Normal"]))
            story.append(Spacer(1, 0.3 * cm))

            header = ["Data Retirada", "Hora", "Funcionário", "Ferramenta", "Qtd.", "Local", "Devolvido?", "Data Devolução"]
            data = [header]

            for row in results:
                key = (row.codigo_item, row.matricula)
                devolvido = "Não"
                data_devolucao = ""
                if key in devolucoes_map and devolucoes_map[key]:
                    devolvido = "Sim"
                    data_devolucao = TimeService.format_local(devolucoes_map[key][0]["data"], "%d/%m/%Y")

                local_info = row.local_servico or "N/D"
                if row.observacao and row.observacao.strip():
                    local_info = f"{local_info} | {row.observacao}"

                data.append([
                    TimeService.format_local(row.data_saida, "%d/%m/%Y"),
                    TimeService.format_local(row.data_saida, "%H:%M"),
                    Paragraph((row.usuario_nome or "N/D")[:25], body_style),
                    Paragraph((row.item_descricao or "Ferramenta removida")[:35], body_style),
                    row.quantidade,
                    Paragraph(local_info[:150], body_style),
                    devolvido,
                    data_devolucao,
                ])

            table = Table(
                data,
                colWidths=[2.2 * cm, 1.3 * cm, 5.0 * cm, 6.5 * cm, 1.3 * cm, 8.0 * cm, 1.8 * cm, 2.5 * cm],
            )
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 0), (-1, -1), 7),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("ALIGN", (0, 0), (-1, 0), "LEFT"),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f6")]),
                        ("LEFTPADDING", (0, 0), (-1, -1), 3),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                        ("TOPPADDING", (0, 0), (-1, -1), 2),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                    ]
                )
            )
            story.append(table)
            doc.build(story)

            buffer.seek(0)
            return send_file(
                buffer,
                as_attachment=True,
                download_name=pdf_filename,
                mimetype="application/pdf",
            )
    
    except Exception as e:
        logger.exception("Erro ao gerar relatório de ferramentas")
        flash(f"Erro ao gerar relatório: {str(e)}", "danger")
        return _redirect_back()


@bp.route("/relatorios/perdidas-quebradas.xlsx")
@login_required
def lost_broken_report_xlsx():
    """Gera planilha A4 vertical de ferramentas perdidas e quebradas."""
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from ..utils.report_branding import get_company_header_lines
    from ..utils.time_service import TimeService

    rows = tool_custody_service.list_lost_broken_tools(limit=None)

    wb = Workbook()
    ws = wb.active
    ws.title = "Perdidas e Quebradas"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.freeze_panes = "A8"

    total_columns = 12
    title_fill = PatternFill("solid", fgColor="0F172A")
    subtitle_fill = PatternFill("solid", fgColor="1E293B")
    header_fill = PatternFill("solid", fgColor="2563EB")
    group_fill = PatternFill("solid", fgColor="E2E8F0")
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_columns)
    ws["A1"] = "Relatório de Ferramentas"
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_columns)
    ws["A2"] = "Perdidas e quebradas"
    ws["A2"].font = Font(color="FFFFFF", bold=True, size=12)
    ws["A2"].fill = subtitle_fill
    ws["A2"].alignment = Alignment(horizontal="center")

    current_row = 3
    for line in get_company_header_lines():
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=total_columns)
        ws.cell(current_row, 1, line)
        ws.cell(current_row, 1).alignment = Alignment(horizontal="center")
        current_row += 1

    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=total_columns)
    ws.cell(current_row, 1, f"Gerado em {TimeService.now_local().strftime('%d/%m/%Y %H:%M')} | Total: {len(rows)} ocorrência(s)")
    ws.cell(current_row, 1).alignment = Alignment(horizontal="center")
    current_row += 2

    headers = [
        "Colaborador", "Matrícula", "Situação", "Data retirada", "Data ocorrência", "Local/uso",
        "Código", "Ferramenta", "Marca", "Qtd.", "Valor ref.", "Motivo / situação",
    ]
    header_row = current_row
    ws.append(headers)
    for cell in ws[header_row]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    current_row += 1

    current_employee = None
    for row in rows:
        employee_key = f"{row.get('colaborador') or 'N/D'} ({row.get('matricula') or 'N/D'})"
        if employee_key != current_employee:
            current_employee = employee_key
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=total_columns)
            group_cell = ws.cell(current_row, 1, current_employee)
            group_cell.fill = group_fill
            group_cell.font = Font(color="0F172A", bold=True)
            group_cell.alignment = Alignment(horizontal="left")
            current_row += 1

        valor_total = row.get("valor_total")
        valor_label = f"{float(valor_total):.2f} ({row.get('valor_fonte') or 'valor ref.'})" if valor_total is not None else ""
        ws.append([
            row.get("colaborador") or "N/D",
            row.get("matricula") or "N/D",
            row.get("kind_label") or "N/D",
            row.get("data_retirada_label") or "N/D",
            row.get("data_evento_label") or "N/D",
            row.get("local_servico") or "Não informado",
            row.get("codigo_item") or "N/D",
            row.get("descricao") or "Ferramenta",
            row.get("marca") or "N/D",
            row.get("quantidade") or 0,
            valor_label,
            row.get("motivo") or "Motivo não informado",
        ])
        for cell in ws[current_row]:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        current_row += 1

    widths = [24, 14, 12, 18, 18, 24, 16, 34, 16, 8, 18, 42]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"relatorio_ferramentas_perdidas_quebradas_{TimeService.now_local().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@bp.route("/api/funcionarios/buscar")
@login_required
def search_employees():
    """Busca funcionários por nome ou matrícula (API)."""
    from ..models import Usuario
    
    query = request.args.get("q", "").strip()
    
    if not query or len(query) < 1:
        return jsonify([])
    
    # Buscar usuários que correspondem à query
    usuarios = (
        Usuario.query.filter(
            or_(
                Usuario.nome.ilike(f"%{query}%"),
                Usuario.matricula.ilike(f"%{query}%")
            )
        )
        .limit(20)
        .all()
    )
    
    return jsonify([{
        "matricula": u.matricula,
        "nome": u.nome,
        "cargo": u.cargo or "N/D",
        "setor": u.setor or "N/D"
    } for u in usuarios])


@bp.route("/api/ferramentas/buscar")
@login_required
def search_tools():
    """Busca ferramentas por código, descrição ou marca (API)."""
    from ..services.inventory import inventory_service

    query = request.args.get("q", "").strip()
    if not query or len(query) < 1:
        return jsonify([])

    resultados = inventory_service.search_items_for_autocomplete(query, limit=20)
    ferramentas = [
        item for item in resultados
        if "ferrament" in str(item.get("categoria") or "").lower()
        and item.get("is_available") is not False
    ]

    return jsonify([
        {
            "codigo": item.get("codigo") or "",
            "descricao": item.get("descricao") or "",
            "categoria": item.get("categoria") or "N/D",
            "marca": item.get("marca") or "N/D",
            "saldo": item.get("saldo") or 0,
            "saldo_display": item.get("saldo_display") or "0",
            "saldo_disponivel": item.get("saldo_disponivel") or 0,
            "saldo_disponivel_display": item.get("saldo_disponivel_display") or "0",
            "unidade": item.get("unidade") or "un.",
            "is_available": item.get("is_available"),
            "unavailable_reason": item.get("unavailable_reason"),
            "unavailable_detail": item.get("unavailable_detail"),
            "has_open_repair": item.get("has_open_repair"),
        }
        for item in ferramentas
    ])


@bp.route("/api/funcionario/<matricula>/historico")
@login_required
def get_employee_full_history(matricula: str):
    """Retorna histórico completo de materiais E ferramentas do funcionário (API)."""
    from ..models import Saida, Item, Usuario, InventarioEvento
    from ..utils.time_service import TimeService
    
    # Verificar se funcionário existe
    usuario = Usuario.query.filter_by(matricula=matricula).first()
    if not usuario:
        return jsonify({"success": False, "error": "Funcionário não encontrado"}), 404
    
    # Buscar TODAS as saídas do funcionário (sem limite de tempo)
    saidas = (
        db.session.query(Saida, Item)
        .join(Item, Saida.codigo_item == Item.codigo_item)
        .filter(Saida.matricula == matricula)
        .order_by(Saida.data_saida.desc())
        .all()
    )
    
    # Separar por categoria
    materiais = []
    ferramentas = []
    
    for saida, item in saidas:
        data_formatada = TimeService.format_local(saida.data_saida, "%d/%m/%Y")
        hora = TimeService.format_local(saida.data_saida, "%H:%M")
        
        # Para ferramentas, verificar se foi devolvida
        status_devolucao = ""
        categoria_lower = (item.categoria or "").lower()
        if "ferrament" in categoria_lower:
            # Buscar devolução no InventarioEvento
            devolucao = (
                db.session.query(InventarioEvento)
                .filter(
                    InventarioEvento.matricula == matricula,
                    InventarioEvento.codigo_item == item.codigo_item,
                    InventarioEvento.tipo.in_(["devolucao_ferramenta", "devolucao"]),
                    InventarioEvento.data_evento >= saida.data_saida
                )
                .first()
            )
            status_devolucao = "✅ Devolvido ao estoque" if devolucao else "⚠️ Pendência !!"
        
        data = {
            "id": saida.id_saida,
            "data_formatada": data_formatada,
            "hora": hora,
            "item": item.descricao or "Item removido",
            "codigo": item.codigo_item or "N/D",
            "quantidade": saida.quantidade or 0,
            "local": saida.local_servico or "",
            "observacao": saida.observacao or "",
            "status_devolucao": status_devolucao
        }
        
        # Classificar por categoria
        if "ferrament" in categoria_lower:
            ferramentas.append(data)
        else:
            materiais.append(data)
    
    return jsonify({
        "success": True,
        "usuario": {
            "nome": usuario.nome,
            "matricula": usuario.matricula,
            "cargo": usuario.cargo or "N/D"
        },
        "materiais": materiais,
        "ferramentas": ferramentas
    })


@bp.route("/api/funcionario/<matricula>/relatorio.pdf")
@login_required
def download_employee_report(matricula: str):
    """Gera e baixa relatório (PDF) do histórico de um funcionário.

    Querystring:
      - aba: "materiais" | "ferramentas" | "all" (padrão: all)
    """
    from io import BytesIO

    from ..models import Saida, Usuario
    from ..utils.time_service import TimeService
    from ..views.reports import _generate_usuario_report_pdf

    aba = (request.args.get("aba") or "all").strip().lower()
    if aba not in {"materiais", "ferramentas", "all"}:
        aba = "all"
    
    # Verificar se funcionário existe
    usuario = Usuario.query.filter_by(matricula=matricula).first()
    if not usuario:
        flash("Funcionário não encontrado", "danger")
        return redirect(url_for("tool_custody.index"))
    
    # Buscar todas as saídas
    saidas_query = (
        db.session.query(Saida)
        .filter(Saida.matricula == matricula)
        .order_by(Saida.data_saida.desc())
        .all()
    )
    
    # Converter para formato compatível com o gerador de PDF
    saidas_data = []
    for saida in saidas_query:
        categoria_lower = ""
        if getattr(saida, "item", None) and getattr(saida.item, "categoria", None):
            categoria_lower = (saida.item.categoria or "").lower()

        is_ferramenta = "ferrament" in categoria_lower
        if aba == "ferramentas" and not is_ferramenta:
            continue
        if aba == "materiais" and is_ferramenta:
            continue

        saidas_data.append({
            "data": saida.data_saida,
            "item_descricao": saida.item.descricao if saida.item else "Item removido",
            "codigo_item": saida.codigo_item,
            "quantidade": saida.quantidade,
            "periodo": TimeService.get_business_day_tag(saida.data_saida),
            "observacao": saida.observacao,
            "local_servico": saida.local_servico,
            "tipo": "Retirada"
        })
    
    usuario_data = {
        "nome": usuario.nome,
        "matricula": usuario.matricula,
        "cargo": usuario.cargo or "N/D"
    }
    
    try:
        # Gerar PDF usando a função existente
        pdf_bytes = _generate_usuario_report_pdf(
            usuario=usuario_data,
            saidas=saidas_data,
            period_days=0,  # 0 = todo o histórico
            time_service=TimeService
        )
        
        # Retornar o arquivo PDF
        return send_file(
            BytesIO(pdf_bytes),
            mimetype="application/pdf",
            as_attachment=True,
            download_name=f"relatorio_{usuario.nome.replace(' ', '_')}_{matricula}.pdf"
        )
    except Exception as e:
        flash(f"Erro ao gerar relatório: {str(e)}", "danger")
        return redirect(url_for("tool_custody.index"))
