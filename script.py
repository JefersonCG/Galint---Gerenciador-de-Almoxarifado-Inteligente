import openpyxl

class PurchaseProjectionService:
    @staticmethod
    def build_workbook(report_data, filename):
        wb = openpyxl.Workbook()
        ws = wb.active
        headers = ["Item", "Supplier Contact", "Product Link", "Unit Price", "Total Price", "Link Fornecedor"]
        ws.append(headers)
        
        for category in report_data.get("categories", []):
            for item in category.get("items", []):
                quote = item.get("selected_potential_quote", {})
                ws.append([
                    item.get("name"),
                    quote.get("contact_url"),
                    quote.get("product_url"),
                    quote.get("unit_price"),
                    quote.get("total_price"),
                    quote.get("contact_url")  # Link Fornecedor
                ])
        wb.save(filename)

report = {
    "categories": [{
        "items": [{
            "name": "Test Item",
            "selected_potential_quote": {
                "contact_url": "http://supplier.com/contact",
                "product_url": "http://supplier.com/product",
                "unit_price": 10.0,
                "total_price": 100.0
            }
        }]
    }]
}

filename = "projection.xlsx"
PurchaseProjectionService.build_workbook(report, filename)

wb = openpyxl.load_workbook(filename)
ws = wb.active
header_row = [cell.value for cell in ws[1]]
print(f"Header Row: {header_row}")
print(f"Link Fornecedor Present: {'Link Fornecedor' in header_row}")
print(f"Total Header Count: {len(header_row)}")
